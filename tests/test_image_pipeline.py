import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from PIL import Image
import pymupdf

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'scripts'))
from report_images import attach_question_images, item_tables, keep_item_analysis_connected
from qa_verify import verify_question_images
from generate_reports import build_question_dict, generate_chapter_4
from source_inputs import missing_statistics
from question_screenshots import page_bounds, exclude_question_number, exclude_exam_transition, exclude_section_heading, render


class ImagePipelineTests(unittest.TestCase):
    def test_next_section_heading_is_excluded_but_all_options_remain(self):
        with pymupdf.open() as source:
            page=source.new_page()
            page.insert_text((50,70),'12. Question')
            for i,letter in enumerate('ABCD'):
                page.insert_text((60,100+25*i),letter+'. Option')
            page.insert_text((50,230),'二、閱讀理解題',fontname='china-t')
            record={'kind':'question','parts':[{'page':1,'rect':[30,40,500,260]}]}
            exclude_section_heading(source,record)
            text=page.get_text(clip=pymupdf.Rect(record['parts'][0]['rect']))
            self.assertNotIn('閱讀理解題',text)
            for letter in 'ABCD':self.assertIn(letter+'. Option',text)
            self.assertIn('excluded_section_heading',record['parts'][0])

    def test_section_heading_inside_article_frame_is_preserved(self):
        with pymupdf.open() as source:
            page=source.new_page()
            page.draw_rect(pymupdf.Rect(40,50,400,200))
            page.insert_text((50,90),'二、閱讀理解題',fontname='china-t')
            record={'kind':'question','parts':[{'page':1,'rect':[30,40,500,260]}]}
            exclude_section_heading(source,record)
            self.assertEqual(record['parts'][0]['rect'][3],260)

    def test_chinese_layout_preserves_original_content(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'第四章_C7N_AI產出.docx'
            analysis=missing_statistics(['C','B'])
            template=Path(__file__).resolve().parents[1]/'resources/templates/pydocxTemplet_C_new.docx'
            generate_chapter_4('C','7N','115',['C','B'],[('A','形音知識'),('G','字詞理解')],analysis,template,path)
            doc=Document(path)
            tables=item_tables(doc)
            self.assertEqual(tables[1].cell(1,len(tables[1].columns)-1).text,'C')
            self.assertEqual(tables[1].cell(3,1).text,'')
            self.assertEqual(tables[2].cell(3,1).text,'')
            self.assertIn('本題主要檢測學生形音知識的能力。',tables[1].cell(9,0).text)
            text=''.join(p.text for p in doc.paragraphs)
            self.assertIn('七年級評量的語文知識',text)
            self.assertIn('七年級評量的閱讀理解',text)
            intros=[p for p in doc.paragraphs if '年級評量的' in p.text]
            self.assertEqual(len(intros),2)
            for p in intros:
                ind=p._p.xpath('./w:pPr/w:ind')[0]
                from docx.oxml.ns import qn
                self.assertEqual(ind.get(qn('w:left')),'840')
                self.assertEqual(ind.get(qn('w:firstLine')),'480')
                node=p._p.getnext()
                while node is not None and node.tag!=qn('w:tbl'):
                    self.assertFalse(node.xpath('.//w:br[@w:type="page"]'))
                    node=node.getnext()
                self.assertIsNotNone(node)

    def test_listening_transition_is_outside_question_crop(self):
        with pymupdf.open() as source:
            page=source.new_page()
            page.insert_text((50,70),'15. A. They watch TV at home.')
            page.insert_text((50,100),'D. They play basketball at the park.')
            page.insert_text((50,160),'各位同學，聽力試題已播音完畢，下一大題開始為閱讀測驗。',fontname='china-t')
            page.insert_text((50,185),'現在請翻到下一頁繼續完成作答。',fontname='china-t')
            record={'kind':'question','number':15,'parts':[{'page':1,'rect':[30,40,550,220]}]}
            exclude_exam_transition(source,record)
            clip=pymupdf.Rect(record['parts'][0]['rect'])
            text=page.get_text(clip=clip)
            self.assertIn('D. They play basketball',text)
            self.assertNotIn('聽力試題',text)
            self.assertNotIn('現在請翻',text)

    def test_statistics_connect_to_analysis_without_locking_entire_row(self):
        document=Document()
        table=document.add_table(rows=10,cols=3)
        table.cell(0,0).text='題號'
        table.cell(1,0).text='1'
        cell=table.cell(9,0)
        cell.text='一、試題分析'
        cell.add_paragraph('(一) 評量目標')
        cell.add_paragraph('目標內容')
        cell.add_paragraph('(二) 數據說明')
        cell.add_paragraph('數據內容')
        keep_item_analysis_connected(document)
        self.assertFalse(table.rows[9]._tr.xpath('./w:trPr/w:cantSplit'))
        self.assertFalse(table.cell(8,0).paragraphs[0].paragraph_format.keep_with_next)
        self.assertIn('目標內容',table.cell(9,0).text)
        before=table.cell(9,0).text
        breaks_before=len(document.element.xpath('.//w:br[@w:type="page"]'))
        keep_item_analysis_connected(document)
        self.assertEqual(table.cell(9,0).text,before)
        self.assertEqual(len(table.rows),10)
        self.assertEqual(len(document.element.xpath('.//w:br[@w:type="page"]')),breaks_before)
        from docx.shared import Pt
        from docx.enum.text import WD_LINE_SPACING
        for row in table.rows[9:]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    self.assertEqual(paragraph.paragraph_format.line_spacing,Pt(24))
                    self.assertEqual(paragraph.paragraph_format.line_spacing_rule,WD_LINE_SPACING.EXACTLY)
                    self.assertFalse(paragraph.paragraph_format.keep_together)

    def test_only_question_number_is_removed_from_same_line(self):
        with pymupdf.open() as source:
            page=source.new_page()
            page.insert_text((50,70),'5. Calculate 5.25 + 5.75 = ?')
            page.insert_text((50,100),'A. 10   B. 11')
            record={'kind':'question','number':5,'parts':[{'page':1,'rect':[30,40,350,130]}]}
            exclude_question_number(source,record)
            self.assertTrue(record['question_number_removed'])
            box=pymupdf.Rect(record['parts'][0]['exclude_rects'][0])
            self.assertEqual(page.get_text(clip=box).strip(),'5.')
            self.assertLess(box.x1,page.search_for('Calculate')[0].x0)
            image=render(source,record['parts'],144)
            self.assertGreater(image.width,100)
            self.assertIn('5.25',page.get_text())

    def test_fraction_near_page_bottom_is_not_a_page_number(self):
        with pymupdf.open() as source:
            page=source.new_page(width=595,height=842)
            page.insert_text((135,735),'1')
            page.insert_text((297,790),'2')
            _,bottom=page_bounds(page)
            self.assertGreater(bottom,760)
            self.assertLess(bottom,790)

    def fixture(self, root):
        pdf = root/'booklet.pdf'
        with pymupdf.open() as source:
            for label in ['passage page one','passage page two']:
                source.new_page().insert_text((50,70),label)
            source.save(pdf)
        records = []
        for number, color in [(1,'red'),(2,'blue')]:
            filename = f'q{number}.png'
            Image.new('RGB',(200,100),color).save(root/filename)
            records.append({'id':f'q{number}','number':number,'kind':'question','image':filename,
                'parts':[{'page':number,'rect':[30,40,300,150]}],'materials':['group_1-2'],'warnings':[]})
        records.append({'id':'group_1-2','number':None,'kind':'material','image':'unused.png',
            'question_range':[1,2],'parts':[{'page':n,'rect':[30,40,300,150]} for n in [1,2]],'warnings':[]})
        manifest = root/'manifest.json'
        manifest.write_text(json.dumps({'source_pdf':str(pdf),'dpi':120,'records':records}),encoding='utf-8')
        doc = Document()
        for number in [2,1]:  # 報告向度排序與題號順序不同。
            table = doc.add_table(rows=10,cols=3)
            table.cell(0,0).text='題號'
            table.cell(1,0).text=str(number)
            table.cell(1,1).text=f'115-C6-A1-{number:02d}'
            table.cell(4,0).text='題目'
            table.cell(9,0).text='專家分析待填'
        report = root/'第四章_C7N_AI產出.docx'
        doc.save(report)
        return report,manifest

    def test_actual_numbers_hashes_and_passage_inserted_once(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            report,manifest=self.fixture(root)
            attach_question_images(report,manifest)
            checks=verify_question_images(report)
            self.assertFalse(any(c[1]=='FAIL' for c in checks),checks)
            doc=Document(report)
            self.assertEqual(list(item_tables(doc)),[2,1])
            descriptions=[n.get('descr','') for n in doc.element.xpath('.//wp:docPr')]
            self.assertEqual(sum(d.startswith('question:') for d in descriptions),2)
            self.assertEqual(descriptions.count('material:group_1-2:combined'),1)
            self.assertLess(descriptions.index('material:group_1-2:combined'),next(i for i,d in enumerate(descriptions) if d.startswith('question:q2;')))
            audit=json.loads(report.with_suffix('.images.json').read_text(encoding='utf-8'))
            self.assertEqual(audit['material_layouts'][0]['part_count'],2)
            self.assertTrue(all(t.cell(9,0).text=='專家分析待填' for t in item_tables(doc).values()))

    def test_forced_centering_fails_left_top_policy(self):
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        with tempfile.TemporaryDirectory() as directory:
            report, manifest = self.fixture(Path(directory))
            attach_question_images(report, manifest)
            document = Document(report)
            cell = item_tables(document)[1].cell(4,1)
            self.assertEqual(cell.vertical_alignment, WD_CELL_VERTICAL_ALIGNMENT.TOP)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                if paragraph._p.xpath('.//w:drawing'):
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            document.save(report)
            self.assertTrue(any(c[1]=='FAIL' for c in verify_question_images(report)))

    def test_shared_passage_is_vertically_centered_while_question_stays_top(self):
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        with tempfile.TemporaryDirectory() as directory:
            report,manifest=self.fixture(Path(directory))
            data=json.loads(manifest.read_text(encoding='utf-8'))
            data['identity']='115_桃園市_國語文_7年級'
            manifest.write_text(json.dumps(data),encoding='utf-8')
            attach_question_images(report,manifest)
            document=Document(report)
            self.assertEqual(item_tables(document)[1].cell(4,1).vertical_alignment,WD_CELL_VERTICAL_ALIGNMENT.TOP)
            articles=[t for t in document.tables if len(t.rows)==1 and t._tbl.xpath('.//wp:docPr[starts-with(@descr,"material:")]')]
            self.assertEqual(len(articles),1)
            self.assertEqual(articles[0].cell(0,0).vertical_alignment,WD_CELL_VERTICAL_ALIGNMENT.CENTER)
            self.assertFalse(any(c[1]=='FAIL' for c in verify_question_images(report)))
            articles[0].cell(0,0).vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            document.save(report)
            self.assertTrue(any(c[1]=='FAIL' for c in verify_question_images(report)))

    def test_missing_image_does_not_change_report(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            report,manifest=self.fixture(root)
            before=report.read_bytes()
            (root/'q2.png').unlink()
            with self.assertRaises(ValueError):
                attach_question_images(report,manifest)
            self.assertEqual(report.read_bytes(),before)

    def test_short_and_long_questions_use_same_source_scale(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            report,manifest=self.fixture(root)
            Image.new('RGB',(800,500),'blue').save(root/'q2.png')
            attach_question_images(report,manifest)
            document=Document(report)
            extents={number:table.cell(4,1)._tc.xpath('.//wp:extent')[0]
                     for number,table in item_tables(document).items()}
            self.assertAlmostEqual(int(extents[2].get('cx'))/int(extents[1].get('cx')),4,places=4)
            extents[1].set('cx',str(int(extents[1].get('cx'))*2))
            document.save(report)
            self.assertTrue(any(c[1]=='FAIL' and 'scale' in c[0] for c in verify_question_images(report)))

    def test_range_only_material_stays_inside_first_question(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            report,manifest=self.fixture(root)
            with pymupdf.open() as source:
                source.new_page().insert_text((50,70),'1-2')
                source.new_page().insert_text((50,70),'question two')
                source.save(root/'range.pdf')
            data=json.loads(manifest.read_text(encoding='utf-8'))
            data['source_pdf']=str(root/'range.pdf')
            data['records'][-1]['parts']=[{'page':1,'rect':[30,40,300,150]}]
            manifest.write_text(json.dumps(data),encoding='utf-8')
            attach_question_images(report,manifest)
            doc=Document(report)
            self.assertEqual(len(item_tables(doc)[2].cell(4,1)._tc.xpath('.//w:drawing')),2)
            self.assertFalse(any('共用文章' in p.text for p in doc.paragraphs))
            self.assertFalse(any(c[1]=='FAIL' for c in verify_question_images(report)))

    def test_duplicate_or_missing_number_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            report,manifest=self.fixture(root)
            data=json.loads(manifest.read_text(encoding='utf-8'))
            data['records']=[r for r in data['records'] if r['id']!='q2']
            manifest.write_text(json.dumps(data),encoding='utf-8')
            with self.assertRaises(ValueError):
                attach_question_images(report,manifest)

    def test_changed_embedded_png_fails_qa(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            report,manifest=self.fixture(root)
            attach_question_images(report,manifest)
            document=Document(report)
            blip=item_tables(document)[1].cell(4,1)._tc.xpath('.//a:blip')[0]
            from docx.oxml.ns import qn
            part=document.part.related_parts[blip.get(qn('r:embed'))]
            member=str(part.partname).lstrip('/')
            replacement=root/'replacement.png'
            Image.new('RGB',(200,100),'green').save(replacement)
            with ZipFile(report) as source:
                entries={name:source.read(name) for name in source.namelist()}
            entries[member]=replacement.read_bytes()
            with ZipFile(report,'w',ZIP_DEFLATED) as output:
                for name,data in entries.items():output.writestr(name,data)
            self.assertTrue(any(c[1]=='FAIL' for c in verify_question_images(report)))

    def test_wrong_subject_or_grade_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            report,manifest=self.fixture(root)
            before=report.read_bytes()
            for identity in ['115_桃園市_數學_7年級','115_國語文_5年級','114_國語文_7年級']:
                data=json.loads(manifest.read_text(encoding='utf-8'))
                data['identity']=identity
                manifest.write_text(json.dumps(data),encoding='utf-8')
                with self.assertRaises(ValueError):
                    attach_question_images(report,manifest)
                self.assertEqual(report.read_bytes(),before)

    def test_draft_never_uses_zero_as_observed_statistics(self):
        analysis=missing_statistics(['A'])
        question=build_question_dict(analysis['questions'][0],'115-C6-A1-01',1,'A','形音知識')
        for key in ['pass_rate','discrimination','hi_pass','lo_pass','all_1','all_9','w1_rate']:
            self.assertEqual(question[key],'待補')
        self.assertEqual(question['ans'],'A')


if __name__=='__main__':
    unittest.main()
