import sys
from pathlib import Path
import unittest
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from proper_names import underline_counties


class ProperNameTests(unittest.TestCase):
    def test_county_across_runs_preserves_text_and_other_formatting(self):
        doc=Document();p=doc.add_paragraph('115年')
        p.add_run('桃').bold=True
        p.add_run('園市國語文')
        before=p.text
        underline_counties(doc)
        self.assertEqual(p.text,before)
        self.assertEqual(''.join(r.text for r in p.runs if r.underline),'桃園市')
        self.assertTrue(next(r for r in p.runs if r.text=='桃').bold)
        underline_counties(doc)
        self.assertEqual(p.text,before)
        self.assertEqual(''.join(r.text for r in p.runs if r.underline),'桃園市')

    def test_generic_county_word_and_other_names_are_not_underlined(self):
        doc=Document();p=doc.add_paragraph('各縣市成果報告：桃園市、新竹縣。')
        underline_counties(doc)
        self.assertEqual(''.join(r.text for r in p.runs if r.underline),'桃園市新竹縣')
