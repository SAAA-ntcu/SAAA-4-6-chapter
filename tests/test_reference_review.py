from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document
from PIL import Image

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from inspect_reference_docs import inspect,scan
from verify_editorial import verify,pending


class ReferenceReviewTests(unittest.TestCase):
    def test_merged_table_does_not_hide_question_or_borrow_next_analysis(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'reference.docx'
            doc=Document();table=doc.add_table(rows=20,cols=3)
            for start,number in [(0,1),(10,2)]:
                for column,label in enumerate(['題號','試題代碼','答案']):table.cell(start,column).text=label
                for column,value in enumerate([str(number),f'115-E6-A1-{number:02d}','A']):table.cell(start+1,column).text=value
                table.cell(start+2,0).text='學習重點';table.cell(start+2,1).text='正式指標內容'
                table.cell(start+9,0).merge(table.cell(start+9,2))
            table.cell(9,0).text='一、試題分析\n(一)評量目標\n本題為……\n(二)數據說明\n統計內容\n(三)選項及學生表現說明\n顯示……\n二、對應學習重點\n三、教學建議\n(一)\n(二)'
            table.cell(19,0).text='一、試題分析\n(一)評量目標\n辨識目標語音\n(二)數據說明\n統計內容\n(三)選項及學生表現說明\n選項差異分析\n二、對應學習重點\n正式對應指標\n三、教學建議\n具體教學步驟'
            doc.save(path)
            self.assertEqual(inspect(path,Path(directory))['item_tables'],2)
            result=verify(path)
            self.assertEqual(result['questions'],2)
            self.assertEqual(result['pending_questions'],1)
            self.assertIn('assessment_target',result['items'][0]['pending_fields'])
            self.assertTrue(result['items'][1]['complete'])
            self.assertFalse(result['editorial_ready'])

    def test_image_names_metadata_and_missing_dpi_are_distinguished(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            for name in ['q1.png','第_02_題.png','題組_1-2.png']:
                Image.new('RGB',(50,30),'white').save(root/name)
            (root/'._q1.png').write_bytes(b'not a PNG')
            (root/'.DS_Store').write_bytes(b'metadata')
            result=scan(root)
            self.assertFalse(result['errors'])
            self.assertEqual(result['ignored_metadata_files'],2)
            images=result['image_sets'][0]
            self.assertEqual(sorted(i['number'] for i in images['questions']),[1,2])
            self.assertEqual(len(images['materials']),1)
            self.assertEqual(images['unknown_dpi'],3)
            self.assertEqual(images['missing_numbers'],[])


if __name__=='__main__':unittest.main()
