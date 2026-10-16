"""唯讀盤點多版本 Word 參考文件；版本階段是線索，不自動選定最終稿。"""
import argparse
from collections import Counter
import hashlib
import html
import json
from pathlib import Path
import re
from zipfile import ZipFile

from lxml import etree
from PIL import Image
import openpyxl

NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
W='{'+NS['w']+'}'


def text(node):
    return ''.join(node.xpath('.//w:t/text()',namespaces=NS)).strip()


def question_blocks(table):
    """Word 會合併相鄰題表；以重複的題號／代碼／答案表頭切回邏輯題目。"""
    rows=table.findall(W+'tr')
    starts=[]
    for index,row in enumerate(rows):
        cells=row.findall(W+'tc')
        labels=[re.sub(r'\s+','',text(c)) for c in cells]
        if len(labels)>=3 and labels[:2]==['題號','試題代碼'] and labels[-1]=='答案':
            starts.append(index)
    for position,start in enumerate(starts):
        end=starts[position+1] if position+1<len(starts) else len(rows)
        block=rows[start:end]
        if len(block)<2:continue
        cells=block[1].findall(W+'tc')
        if len(cells)<3:continue
        code=text(cells[1])
        if not re.fullmatch(r'\d{3}-[CMES]\d-[A-Z]\d-\d{2}',code):continue
        yield start,block,{'number':text(cells[0]),'code':code,'answer':text(cells[-1])}


def inspect(path, root):
    with ZipFile(path) as archive:
        doc=etree.fromstring(archive.read('word/document.xml'))
        tables=doc.xpath('./w:body/w:tbl',namespaces=NS)
        items=[]
        for index,table in enumerate(tables):
            for row_start,block,metadata in question_blocks(table):
                content=''.join(text(row) for row in block)
                analysis=content.find('試題分析')
                spacings=Counter()
                if analysis>=0:
                    active=False
                    for row in block:
                        for p in row.xpath('.//w:p',namespaces=NS):
                            if '試題分析' in text(p):active=True
                            if active and text(p):
                                for s in p.xpath('./w:pPr/w:spacing',namespaces=NS):
                                    spacings[f'{s.get(W+"lineRule","auto")}:{s.get(W+"line","")}']+=1
                items.append({'table_index':index,'row_start':row_start,**metadata,
                    'rows':len(block),
                    'data_text_sha256':hashlib.sha256(''.join(text(row) for row in block[:9]).encode()).hexdigest(),
                    'images':sum(len(row.xpath('.//w:drawing',namespaces=NS)) for row in block),
                    'analysis_line_spacing':dict(spacings),
                    'sample':content[:2200] if not items else None})
        comments=[]
        if 'word/comments.xml' in archive.namelist():
            comments_doc=etree.fromstring(archive.read('word/comments.xml'))
            comments=[text(c)[:700] for c in comments_doc.findall(W+'comment')]
        fonts=Counter(doc.xpath('.//w:rFonts/@w:eastAsia',namespaces=NS))
        sizes=Counter(doc.xpath('.//w:sz/@w:val',namespaces=NS))
        paragraphs=[text(p) for p in doc.xpath('./w:body/w:p',namespaces=NS) if text(p)]
        images=[{'width_cm':round(int(s.get('cx'))/360000,2),
                 'height_cm':round(int(s.get('cy'))/360000,2)}
                for s in doc.xpath('.//wp:extent',namespaces=NS)]
        return {'path':str(path.resolve()),'relative_path':str(path.relative_to(root)),
            'subject_folder':path.relative_to(root).parts[0],
            'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
            'tables':len(tables),'item_tables':len(items),'items':items,
            'body_excerpt':paragraphs[:16],'table_shapes':[
                {'rows':len(t.findall(W+'tr')),'columns':len(t.xpath('./w:tblGrid/w:gridCol',namespaces=NS)),
                 'header':text(t)[:140]} for t in tables[:3]],
            'fonts':dict(fonts),'font_sizes_half_points':dict(sizes),'images':images,
            'insertions':len(doc.xpath('.//w:ins',namespaces=NS)),
            'deletions':len(doc.xpath('.//w:del',namespaces=NS)),
            'comments':comments,'sections':[
                etree.tostring(s,encoding='unicode') for s in doc.xpath('.//w:sectPr',namespaces=NS)]}


def scan(root):
    root=Path(root).resolve()
    reports=[];errors=[]
    for path in sorted(root.rglob('*.docx')):
        if path.name.startswith(('~$','._')):continue
        try:reports.append(inspect(path,root))
        except Exception as exc:errors.append({'path':str(path),'error':str(exc)})
    image_sets={};workbooks=[];ignored=0
    for path in sorted(root.rglob('*')):
        if not path.is_file():continue
        if path.name.startswith(('._','~$')) or path.name=='.DS_Store':
            ignored+=1;continue
        if path.suffix.lower()=='.png':
            key=str(path.parent.relative_to(root))
            group=image_sets.setdefault(key,{'folder':key,'questions':[], 'materials':[], 'other_images':[], 'unknown_dpi':0})
            try:
                with Image.open(path) as image:
                    entry={'file':path.name,'width_px':image.width,'height_px':image.height,
                           'dpi':image.info.get('dpi'),'sha256':hashlib.sha256(path.read_bytes()).hexdigest()}
                    image.verify()
                if not entry['dpi']:group['unknown_dpi']+=1
                number=re.fullmatch(r'(?:q(\d+)|第_0*(\d+)_題)',path.stem)
                if number:
                    entry['number']=int(number[1] or number[2]);group['questions'].append(entry)
                elif path.stem.startswith('題組_'):group['materials'].append(entry)
                else:group['other_images'].append(entry)
            except Exception as exc:errors.append({'path':str(path),'error':str(exc)})
        elif path.suffix.lower()=='.xlsx':
            try:
                book=openpyxl.load_workbook(path,read_only=True,data_only=True)
                workbooks.append({'path':str(path.resolve()),'sheets':[
                    {'name':sheet.title,'rows':sheet.max_row,'columns':sheet.max_column} for sheet in book]})
                book.close()
            except Exception as exc:errors.append({'path':str(path),'error':str(exc)})
    for group in image_sets.values():
        numbers=[i['number'] for i in group['questions']]
        group['missing_numbers']=sorted(set(range(1,max(numbers,default=0)+1))-set(numbers))
        group['duplicate_numbers']=sorted(n for n,c in Counter(numbers).items() if c>1)
    candidates={}
    for report in reports:
        for item in report['items']:
            cohort=item['code'].split('-')[:2]
            key='-'.join(cohort)+':'+item['number']
            candidates.setdefault(key,[]).append({'path':report['path'],'code':item['code'],
                'answer':item['answer'],'data_text_sha256':item['data_text_sha256']})
    conflicts=[{'question':key,'versions':values} for key,values in candidates.items()
               if len({v['data_text_sha256'] for v in values})>1]
    return {'root':str(root),'count':len(reports),'errors':errors,
        'subjects':dict(Counter(r['subject_folder'] for r in reports)),'documents':reports,
        'image_sets':list(image_sets.values()),'workbooks':workbooks,'ignored_metadata_files':ignored,
        'version_conflicts':conflicts}


def write_summary(result,path):
    esc=html.escape
    documents=result['documents'];images=result['image_sets']
    merged=sum(any(i['row_start']>0 for i in d['items']) for d in documents)
    image_count=sum(len(g['questions']) for g in images)
    material_count=sum(len(g['materials']) for g in images)
    unknown=sum(g['unknown_dpi'] for g in images)
    intro=(f'本次唯讀掃描 {len(documents)} 份 Word。{merged} 份含多個邏輯題目合併在同一表格；'
           f'{sum(bool(d["comments"]) for d in documents)} 份含批註，'
           f'{sum(bool(d["insertions"]+d["deletions"]) for d in documents)} 份含文字追蹤修訂。')
    rows=[]
    for d in documents:
        href=Path(d['path']).as_uri()
        rows.append('<tr><td>'+esc(d['subject_folder'])+'</td><td><a href="'+esc(href)+'">'+esc(d['relative_path'])+'</a></td>'
            +''.join('<td>'+str(value)+'</td>' for value in [d['tables'],d['item_tables'],len(d['comments']),d['insertions']+d['deletions']])+'</tr>')
    conflicts=[]
    for group in result['version_conflicts']:
        versions='<br>'.join('<a href="'+esc(Path(v['path']).as_uri())+'">'+esc(str(Path(v['path']).relative_to(result['root'])))+'</a>：'+esc(v['code']) for v in group['versions'])
        conflicts.append('<tr><td>'+esc(group['question'])+'</td><td>'+versions+'</td></tr>')
    image_rows=[]
    for g in images:
        image_rows.append('<tr><td>'+esc(g['folder'])+'</td>'+''.join('<td>'+str(value)+'</td>' for value in [len(g['questions']),len(g['materials']),g['unknown_dpi'],g['missing_numbers'],g['duplicate_numbers']])+'</tr>')
    page='''<!doctype html><html lang="zh-Hant"><meta charset="utf-8"><title>115word 分析與 skill 優化</title>
<style>body{font:17px system-ui;line-height:1.7;margin:36px;max-width:1500px}h1,h2{color:#1e293b}table{border-collapse:collapse;width:100%;margin:15px 0 30px}td,th{padding:10px;text-align:left;border-bottom:1px solid #d8dee8}th{background:#eef2f6}a{color:#185ca3}input{font:inherit;padding:8px;width:80%}li{margin:8px 0}</style>
<h1>115word 分析與 skill 優化</h1><p>'''+intro+'''</p><h2>本次採用的優化</h2><ul>
<li>以重複表頭辨識邏輯題目，避免多題合併後只讀到第一題，並保留來源表格與列位置。</li>
<li>分開檢查資料、版面與專家內容。原稿中的留白及「本題為……」等占位文字應標為待補，不能當作專家內容已完成。</li>
<li>同碼不保證同題：一般六年級與桃園新生不得直接互抄分析與統計。</li>
<li>列出原稿／修訂版的資料差異，明確選定來源，不自動混合不同版本。</li>
<li>兼容 q1 與第_01_題檔名；題組材料分開處理，排除系統附檔。既有圖片沒有 DPI 時不推測物理字級。</li>
<li>保留目前已確認的版面：英文13點等效字級、分析固定行高24點、去題號及排除聽力結束提示。</li></ul>
<p>分析工具不修改原稿，也不接受或拒絕 Word 修訂。完整度通過僅代表欄位有內容，仍需領域審核。</p>
<h2>既有題圖</h2><p>'''+f'{image_count} 張單題圖、{material_count} 張題組圖；{unknown} 張圖片未提供 DPI。忽略 {result["ignored_metadata_files"]} 個系統附檔。'+'''</p>
<table><tr><th>資料夾</th><th>單題圖</th><th>材料圖</th><th>無DPI</th><th>漏號</th><th>重號</th></tr>'''+''.join(image_rows)+'''</table>
<h2>版本差異</h2><p>以下比較每題前9列的文字，包含題碼、向度、答案與統計；不代表圖片內容相同。不同版本的資料需要明確選定來源，不能自動混用。</p>
<table><tr><th>年度科目年級與題號</th><th>候選版本與代碼</th></tr>'''+''.join(conflicts)+'''</table>
<h2>Word 文件清單</h2><input placeholder="輸入科目、年級或檔名篩選" oninput="for(const r of document.querySelectorAll('#docs tbody tr'))r.hidden=!r.textContent.toLowerCase().includes(this.value.toLowerCase())">
<table id="docs"><thead><tr><th>科目</th><th>來源檔案</th><th>實體表格</th><th>邏輯題數</th><th>批註</th><th>文字修訂</th></tr></thead><tbody>'''+''.join(rows)+'''</tbody></table></html>'''
    Path(path).write_text(page,encoding='utf-8')


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--dir',required=True,type=Path)
    parser.add_argument('--json',required=True,type=Path)
    parser.add_argument('--html',type=Path,help='另存可篩選的分析報告')
    args=parser.parse_args()
    result=scan(args.dir)
    args.json.parent.mkdir(parents=True,exist_ok=True)
    args.json.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    if args.html:write_summary(result,args.html)
    print(json.dumps({k:v for k,v in result.items() if k not in ('documents','image_sets','workbooks','version_conflicts')},ensure_ascii=False))
    print('Conflicting question versions:',len(result['version_conflicts']))
    if result['errors']:raise SystemExit(1)
