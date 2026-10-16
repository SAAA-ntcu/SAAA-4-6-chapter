"""以題號讀取原始橫式答案與 PPTX 評量向度，不依檔案排序重新編號。"""
from pathlib import Path
import re

import openpyxl
from pptx import Presentation

CHINESE = {name: chr(65+i) for i, name in enumerate([
    '形音知識', '字詞知識', '語法知識', '修辭知識', '章法知識', '文體知識',
    '字詞理解', '句子理解', '段落理解', '篇章理解'])}
ENGLISH = {'聽力': {'語音聽辨':'A1', '辭彙聽辨':'A2', '教室生活辭句理解與回應':'A3', '文化節慶理解':'A4'},
           '閱讀': {'字詞辨識':'B1', '句子理解':'B2', '短文與對話理解':'B3', '文化節慶理解':'B4'}}


def raw_answers(path):
    result, source_cells = {}, {}
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    for sheet in workbook:
        rows = list(sheet.iter_rows())
        for index, row in enumerate(rows[:-1]):
            if not any(str(c.value).strip() == '題號' for c in row):
                continue
            following = rows[index+1]
            if not any(str(c.value).strip() == '答案' for c in following):
                continue
            for qcell, acell in zip(row, following):
                if not isinstance(qcell.value, (int, float)) or int(qcell.value) != qcell.value:
                    continue
                number, answer = int(qcell.value), str(acell.value or '').strip().upper()
                if answer not in {'A','B','C','D','1','2','3','4'}:
                    raise ValueError(f'答案格式不符：{path.name} {acell.coordinate}={answer}')
                if number in result:
                    raise ValueError(f'答案題號重複：{path.name} q{number}')
                result[number] = answer
                source_cells[number] = f'{sheet.title}!{acell.coordinate}'
    workbook.close()
    if not result or sorted(result) != list(range(1, max(result)+1)):
        raise ValueError(f'找不到完整、連續的橫式答案：{path}')
    return result, source_cells


def raw_dimensions(path, subject, numbers):
    dimensions, sources = {}, {}
    for slide_index, slide in enumerate(Presentation(path).slides, 1):
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            current = ''
            for row_index, row in enumerate(shape.table.rows, 1):
                cells = [re.sub(r'\s+', '', c.text) for c in row.cells]
                if len(cells) < 2 or '題號' in cells[-1]:
                    continue
                if cells[0]:
                    current = cells[0]
                q_numbers = [int(n) for n in re.findall(r'\d+', cells[-1])]
                if not q_numbers:
                    continue
                if subject == 'C':
                    if current not in CHINESE:
                        raise ValueError(f'未定義的國語文向度：{current}')
                    dim = (CHINESE[current], current)
                elif subject == 'E':
                    sub = cells[1]
                    if sub not in ENGLISH.get(current, {}):
                        raise ValueError(f'未定義的英語文向度：{current}/{sub}')
                    dim = (current, ENGLISH[current][sub]+sub)
                else:
                    if cells[1] not in {'概念理解', '程序執行', '解題思考'}:
                        raise ValueError(f'未定義的數學認知向度：{cells[1]}')
                    dim = (current, cells[1])
                for number in q_numbers:
                    if number in dimensions:
                        raise ValueError(f'評量向度題號重複：{path.name} q{number}')
                    dimensions[number] = dim
                    sources[number] = f'投影片{slide_index}/表格第{row_index}列'
    if set(dimensions) != set(numbers):
        raise ValueError(f'向度與答案題號不一致：缺少{sorted(set(numbers)-set(dimensions))}，多出{sorted(set(dimensions)-set(numbers))}')
    return dimensions, sources


def discover_source_sets(input_dir):
    root = Path(input_dir)
    result = []
    for subject, label in [('C','國語文'), ('M','數學'), ('E','英語文')]:
        folder = root/label
        if not folder.is_dir():
            continue
        def exactly_one(paths, kind):
            paths = list(paths)
            if len(paths) != 1:
                raise ValueError(f'{folder} 需要唯一的{kind}，找到 {len(paths)} 份；請整理或改用單科參數')
            return paths[0]
        pdf = exactly_one(folder.glob('*題本.pdf'), '題本 PDF')
        answer = exactly_one(folder.glob('*科答案.xlsx'), '正式答案 Excel')
        dimension = exactly_one(folder.glob('*.pptx'), '評量架構 PPTX')
        answers, answer_cells = raw_answers(answer)
        domains, domain_cells = raw_dimensions(dimension, subject, answers)
        result.append({'subject':subject, 'booklet':pdf, 'answer_file':answer, 'dimension_file':dimension,
            'answers':[answers[n] for n in sorted(answers)], 'domains':[domains[n] for n in sorted(answers)],
            'provenance':{str(n):{'answer':answer_cells[n], 'dimensions':domain_cells[n]} for n in answers}})
    if not result:
        raise ValueError('找不到國語文、數學或英語文科目資料夾')
    return result


def missing_statistics(answers):
    """使用明確的缺值，不能用 0 或示範分數冒充施測結果。"""
    return {'draft':True, 'title':'待補施測統計', 'valid_sample':'待補', 'cronbach':'待補',
        'questions':[{'num':n, 'disc':float('nan'), 'diff':float('nan'), 'pass_rate':float('nan'),
            'correct_ans':answer, 'all_rates':[float('nan')]*5, 'hi_rates':[float('nan')]*5,
            'lo_rates':[float('nan')]*5, 'invalid':False, 'missing_statistics':True}
            for n, answer in enumerate(answers, 1)]}
