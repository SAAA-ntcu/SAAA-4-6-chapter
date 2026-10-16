"""將可追溯的題本 PNG 插入既有第四章表格；文章每題組只插一次。"""
from pathlib import Path
from copy import deepcopy
import hashlib
import json
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph
from PIL import Image
import pymupdf

from question_screenshots import render
from image_layout import plan_question_scale, image_size_cm


def item_tables(document):
    result = {}
    for table in document.tables:
        if len(table.rows) < 5 or table.cell(0, 0).text.strip() != '題號':
            continue
        try:
            number = int(table.cell(1, 0).text.strip())
        except ValueError:
            continue
        if number in result:
            raise ValueError(f'報告題號重複 q{number}')
        result[number] = table
    return result


def add_picture(paragraph, path, width_cm, height_cm, description):
    with Image.open(path) as im:
        width, height = im.size
    final_width = min(width_cm, height_cm*width/height)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.left_indent = Cm(0)
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    shape = paragraph.add_run().add_picture(str(path), width=Cm(final_width))
    shape._inline.docPr.set('descr', description)
    return shape


def article_cell_before(anchor,document,width_cm,height_cm,vertical_center=False):
    """保留單頁文章範圍；國語、英文共用文章靠左且垂直置中。"""
    table=document.add_table(rows=1,cols=1)
    table.autofit=False
    table.alignment=WD_TABLE_ALIGNMENT.LEFT
    table.columns[0].width=Cm(width_cm)
    cell=table.cell(0,0)
    cell.width=Cm(width_cm)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER if vertical_center else WD_CELL_VERTICAL_ALIGNMENT.TOP
    properties=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        border=OxmlElement('w:'+edge);border.set(qn('w:val'),'nil');borders.append(border)
    properties.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge in ['top','left','bottom','right']:
        margin=OxmlElement('w:'+edge);margin.set(qn('w:w'),'0');margin.set(qn('w:type'),'dxa');margins.append(margin)
    properties.append(margins)
    table.rows[0].height=Cm(height_cm)
    table.rows[0].height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    anchor.addprevious(table._tbl)
    return cell


def paragraph_before(anchor, document, text='', page_break=False):
    node = OxmlElement('w:p')
    anchor.addprevious(node)
    paragraph = Paragraph(node, document._body)
    paragraph.paragraph_format.page_break_before = page_break
    run = paragraph.add_run(text)
    run.font.name = '微軟正黑體'
    run.font.size = Pt(12)
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微軟正黑體')
    return paragraph


def passage_anchor(table):
    anchor = table._tbl
    chinese = bool(re.match(r'\d{3}-C\d-', table.cell(1,1).text.strip()))
    previous = anchor.getprevious()
    while previous is not None and previous.tag == qn('w:p'):
        if previous.xpath('.//w:drawing'):
            break
        text = ''.join(previous.itertext()).strip()
        if (not text or re.match(r'^[(（][一二三四五六七八九十]+[)）]', text)
                or re.match(r'^[一二]、', text)
                or (chinese and ('本試題卷評量學生' in text or '年級評量的' in text))):
            anchor, previous = previous, previous.getprevious()
        else:
            break
    return anchor


def normalize_chinese_layout(document):
    """只套用參考稿的引言位置與縮排，不改任何正文或題表內容。"""
    if not any(re.match(r'\d{3}-C\d-',t.cell(1,1).text.strip()) for t in item_tables(document).values()):
        return
    for paragraph in document.paragraphs:
        text=paragraph.text.strip()
        if '年級評量的' in text:
            props=paragraph._p.get_or_add_pPr()
            for old in list(props.findall(qn('w:ind'))):props.remove(old)
            ind=OxmlElement('w:ind')
            ind.set(qn('w:left'),'840')
            ind.set(qn('w:firstLine'),'480')
            props.append(ind)
            paragraph.paragraph_format.line_spacing=Pt(24)
            paragraph.paragraph_format.space_before=Pt(0)
            paragraph.paragraph_format.space_after=Pt(6)
        elif '本試題卷評量學生' in text:
            props=paragraph._p.get_or_add_pPr()
            for old in list(props.findall(qn('w:ind'))):props.remove(old)
            ind=OxmlElement('w:ind')
            ind.set(qn('w:left'),'0');ind.set(qn('w:firstLine'),'480')
            props.append(ind)
            paragraph.paragraph_format.line_spacing=Pt(24)
            paragraph.paragraph_format.space_after=Pt(6)
        elif text in {'一、語文知識：','二、閱讀理解：'}:
            props=paragraph._p.get_or_add_pPr()
            for old in list(props.findall(qn('w:ind'))):props.remove(old)
            ind=OxmlElement('w:ind')
            ind.set(qn('w:left'),'360');ind.set(qn('w:firstLine'),'0')
            props.append(ind)
            paragraph.paragraph_format.line_spacing=Pt(24)
            paragraph.paragraph_format.space_after=Pt(0)


def keep_item_analysis_connected(document):
    """統計表接分析開頭；分析固定行高 24pt，長題允許自然續頁。"""
    normalize_chinese_layout(document)
    def insert_page_break(paragraph):
        boundary=paragraph.insert_paragraph_before()
        boundary.paragraph_format.line_spacing=Pt(1)
        boundary.paragraph_format.space_before=Pt(0)
        boundary.paragraph_format.space_after=Pt(0)
        boundary.paragraph_format.keep_with_next=False
        boundary.paragraph_format.keep_together=False
        boundary.paragraph_format.page_break_before=False
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0')
        boundary._p.get_or_add_pPr().append(snap)
        run=boundary.add_run()
        run.font.size=Pt(1)
        run.add_break(WD_BREAK.PAGE)
    # 清除舊的空白分頁段落，重新建立每題唯一的表外分頁點。
    for paragraph in list(document.paragraphs):
        if (not paragraph.text.strip() and not paragraph._p.xpath('.//w:drawing')
                and not paragraph._p.xpath('.//w:br[@w:type="page"]')):
            following=paragraph._p.getnext()
            if following is not None and following.tag!=qn('w:sectPr'):
                paragraph._p.getparent().remove(paragraph._p)
    for table in item_tables(document).values():
        if len(table.rows)<10:
            continue
        # 相鄰表格在 Word 會自動接成一張；分頁必須放在表格外。
        anchor=passage_anchor(table)
        if anchor.tag==qn('w:p'):
            boundary=Paragraph(anchor,document._body)
        else:
            boundary=paragraph_before(anchor,document)
            boundary.paragraph_format.line_spacing=Pt(1)
            boundary.paragraph_format.space_before=Pt(0)
            boundary.paragraph_format.space_after=Pt(0)
            boundary.add_run().font.size=Pt(1)
        boundary.paragraph_format.page_break_before=False
        boundary.paragraph_format.keep_with_next=False
        # 使用表外分頁符號，避免每個儲存格出現 Word 的黑色方點標記。
        if not boundary._p.xpath('.//w:br[@w:type="page"]'):
            insert_page_break(boundary)
        for node in table._tbl.xpath('.//w:pageBreakBefore'):
            node.getparent().remove(node)
        # 原範本把三個分析章節放在同一個巨大儲存格；Word 會整列移頁。
        # 分成外觀連續的三列，保留文字並移除純空白占位，允許自然換頁。
        paragraphs=[]
        seen_cells=set()
        for row in table.rows[9:]:
            for cell in row.cells:
                if cell._tc in seen_cells:
                    continue
                seen_cells.add(cell._tc)
                paragraphs.extend(deepcopy(p._p) for p in cell.paragraphs if p.text.strip())
        if not paragraphs:
            paragraphs=[deepcopy(table.cell(9,0).paragraphs[0]._p)]
        template=deepcopy(table.rows[9]._tr)
        for row in list(table.rows)[9:]:
            table._tbl.remove(row._tr)
        groups=[[]]
        for node in paragraphs:
            text=''.join(node.itertext()).strip()
            if re.match(r'^[二三]、',text) and groups[-1]:
                groups.append([])
            groups[-1].append(node)
        for group_index,group in enumerate(groups):
            tr=deepcopy(template)
            cells=tr.findall(qn('w:tc'))
            tc=cells[0]
            for extra in cells[1:]:tr.remove(extra)
            tcpr=tc.get_or_add_tcPr()
            span=tcpr.find(qn('w:gridSpan'))
            if span is None:
                span=OxmlElement('w:gridSpan');tcpr.append(span)
            span.set(qn('w:val'),str(len(table.columns)))
            for child in list(tc):
                if child.tag!=qn('w:tcPr'):tc.remove(child)
            for node in group:tc.append(node)
            borders=tcpr.find(qn('w:tcBorders'))
            if borders is not None:
                for edge,visible in [('top',group_index==0),('bottom',group_index==len(groups)-1)]:
                    border=borders.find(qn('w:'+edge))
                    if border is None:
                        border=OxmlElement('w:'+edge);borders.append(border)
                    border.set(qn('w:val'),'single' if visible else 'nil')
                    if visible:border.set(qn('w:sz'),'4')
            properties=tr.get_or_add_trPr()
            for lock in list(properties.findall(qn('w:cantSplit'))):properties.remove(lock)
            table._tbl.append(tr)
        # 各統計列不在列內拆開；避免 keepNext 把整個大型分析列推到次頁。
        seen=set()
        for row_index,row in enumerate(table.rows[:9]):
            row_properties=row._tr.get_or_add_trPr()
            if row_properties.find(qn('w:cantSplit')) is None:
                row_properties.append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next=False
                    paragraph.paragraph_format.keep_together=False
        seen=set()
        for row in table.rows[9:]:
          for cell in row.cells:
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            paragraphs=cell.paragraphs
            for index,paragraph in enumerate(paragraphs):
                text=paragraph.text.strip()
                # 保留字級，設定固定 24pt 行高，取消文件格線額外撐高。
                ppr=paragraph._p.get_or_add_pPr()
                snap=ppr.find(qn('w:snapToGrid'))
                if snap is None:
                    snap=OxmlElement('w:snapToGrid')
                    ppr.append(snap)
                snap.set(qn('w:val'),'0')
                paragraph.paragraph_format.page_break_before=False
                paragraph.paragraph_format.keep_together=False
                paragraph.paragraph_format.keep_with_next=False
                paragraph.paragraph_format.line_spacing=Pt(24)
                paragraph.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY
                paragraph.paragraph_format.space_before=Pt(0)
                paragraph.paragraph_format.space_after=Pt(0)
                if not text:
                    paragraph.paragraph_format.line_spacing=Pt(24)
                    paragraph.paragraph_format.space_before=Pt(0)
                    paragraph.paragraph_format.space_after=Pt(0)
                    for run in paragraph.runs:run.font.size=Pt(1)
                if re.fullmatch(r'[（(][一二三四五六七八九十]+[）)]\s*',text):
                    paragraph.paragraph_format.line_spacing=Pt(24)
                    paragraph.paragraph_format.space_before=Pt(0)
                    paragraph.paragraph_format.space_after=Pt(0)
                for br in paragraph._p.xpath('.//w:br'):
                    trailing=not any((n.text or '').strip() for n in br.xpath('following::w:t')
                                     if n.getparent() is br.getparent())
                    if br.get(qn('w:type'))=='page' or trailing:
                        br.getparent().remove(br)
    # 包含範本正文／文章標題，取消會在編輯標記模式顯示黑方點的段落旗標。
    for paragraph in list(document.paragraphs):
        if paragraph.paragraph_format.page_break_before:
            insert_page_break(paragraph)
        paragraph.paragraph_format.page_break_before=False
        paragraph.paragraph_format.keep_with_next=False
        paragraph.paragraph_format.keep_together=False


def attach_question_images(docx_path, manifest_path):
    manifest_path = Path(manifest_path)
    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
    records = {r['id']:r for r in manifest['records']}
    document = Document(docx_path)
    tables = item_tables(document)
    questions = {r['number']:r for r in records.values() if r['kind']=='question'}
    identity = manifest.get('identity','')
    subject_map = {'國語文':'C','數學':'M','英語文':'E','自然':'S','生物':'S','理化':'S'}
    source_subject = next((code for name,code in subject_map.items() if name in identity),None)
    report_cohort = re.search(r'([CMES])(\d+)(N?)',Path(docx_path).stem)
    for table in tables.values():
        code = re.match(r'(\d{3})-([CMES])',table.cell(1,1).text)
        if code and ((source_subject and code[2]!=source_subject) or (identity[:3].isdigit() and code[1]!=identity[:3])):
            raise ValueError('題本年度／科目與報告試題代碼不符，未插入圖片')
    source_grade = re.search(r'_(\d+)年級$',identity)
    if source_grade and report_cohort and source_grade[1]!=report_cohort[2]:
        raise ValueError('題本施測年級與報告年級不符，未插入圖片')
    if set(tables) != set(questions):
        raise ValueError(f'報告／截圖題號不一致：缺圖{sorted(set(tables)-set(questions))}，多圖{sorted(set(questions)-set(tables))}')
    folder = manifest_path.parent
    for number, record in questions.items():
        if not record.get('image') or not (folder/record['image']).is_file():
            raise ValueError(f'q{number} 截圖不存在')
    inserted_groups, mapping = set(), []
    usable_width = min((s.page_width-s.left_margin-s.right_margin)/360000 for s in document.sections)
    usable_height = min((s.page_height-s.top_margin-s.bottom_margin)/360000 for s in document.sections)
    with pymupdf.open(manifest['source_pdf']) as source:
        widths = [(table.cell(4,1).width/360000 if table.cell(4,1).width else usable_width-2.5)-.35
                  for table in tables.values()]
        layout = plan_question_scale(source, list(questions.values()), folder, manifest['dpi'],
                                     min(14.8,*widths), min(17,usable_height-7),
                                     target_pt=13 if source_subject=='E' else 11)
        material_layouts = []
        for number, table in tables.items():
            record = questions[number]
            cell = table.cell(4, 1)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            if cell._tc.xpath('.//w:drawing') or cell.text.strip():
                raise ValueError(f'q{number} 題目欄已有內容，避免重複插入')
            compact_material = False
            for group_id in record.get('materials', []):
                if group_id in inserted_groups:
                    continue
                material = records[group_id]
                anchor = passage_anchor(table)
                first, last = material.get('question_range', [number,number])
                source_text=''.join(source[p['page']-1].get_text(clip=pymupdf.Rect(p['rect'])).strip()
                                    for p in material['parts'])
                header_only=len(material['parts'])==1 and bool(re.fullmatch(r'[◆◇第題\d\s\-~～至－—]+',source_text))
                image_path=folder/f'{group_id}_整篇_AI產出.png'
                render(source,material['parts'],manifest['dpi']).save(image_path,dpi=(manifest['dpi'],manifest['dpi']))
                natural_width,natural_height=image_size_cm(image_path,manifest['dpi'])
                if header_only:
                    picture=cell.paragraphs[0];compact_material=True
                    material_scale=layout['source_scale']
                    descriptor=f'material:{group_id}:part1'
                else:
                    heading=paragraph_before(anchor,document,f'第{first}～{last}題共用文章／材料',page_break=True)
                    heading.paragraph_format.line_spacing=Pt(18)
                    heading.paragraph_format.space_before=Pt(0)
                    heading.paragraph_format.space_after=Pt(0)
                    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0')
                    heading._p.get_or_add_pPr().append(snap)
                    container=article_cell_before(anchor,document,usable_width,usable_height-1.2,
                                                  vertical_center=source_subject in {'C','E'})
                    picture=container.paragraphs[0]
                    # 整篇文章（含續文）以同一張圖等比例縮放，按單頁可用寬高取最小倍率。
                    material_scale=min((usable_width-.4)/natural_width,(usable_height-1.8)/natural_height)
                    descriptor=f'material:{group_id}:combined'
                shape=add_picture(picture,image_path,natural_width*material_scale,natural_height*material_scale,descriptor)
                material_layouts.append({'id':group_id,'mode':'compact_header' if header_only else 'single_page_composite',
                    'vertical_alignment':'center' if not header_only and source_subject in {'C','E'} else 'top',
                    'target_center_y_cm':document.sections[0].top_margin.cm+.635+(usable_height-1.2)/2
                        if not header_only and source_subject in {'C','E'} else None,
                    'part_count':len(material['parts']),'source_parts':material['parts'],
                    'image':str(image_path.resolve()),'sha256':hashlib.sha256(image_path.read_bytes()).hexdigest(),
                    'question_range':[first,last],'source_scale':material_scale,
                    'effective_body_font_pt':layout['source_body_font_pt']*material_scale if layout['source_body_font_pt'] else None,
                    'width_cm':shape.width/360000,'height_cm':shape.height/360000})
                inserted_groups.add(group_id)
            image_path = folder/record['image']
            paragraph = cell.add_paragraph() if compact_material else cell.paragraphs[0]
            paragraph.clear()
            natural_width, natural_height = layout['sizes'][record['id']]
            shape = add_picture(paragraph, image_path, natural_width*layout['source_scale'], natural_height*layout['source_scale'],
                        f'question:q{number};source:{Path(manifest["source_pdf"]).name};pages:'+','.join(str(p['page']) for p in record['parts']))
            row = table.rows[4]
            row.height = None
            row.height_rule = None
            trPr = row._tr.get_or_add_trPr()
            if trPr.find(qn('w:cantSplit')) is None:
                trPr.append(OxmlElement('w:cantSplit'))
            mapping.append({'question':number, 'code':table.cell(1,1).text,
                'source_pdf':manifest['source_pdf'], 'parts':record['parts'],
                'image':str(image_path.resolve()), 'sha256':hashlib.sha256(image_path.read_bytes()).hexdigest(),
                'display_width_cm':shape.width/360000,'display_height_cm':shape.height/360000,
                'source_scale':layout['source_scale'],
                'materials':record.get('materials',[])})
    # Word 必須在最後表格後放一個段落；若省略，會以預設行高補出空白尾頁。
    tail = document.add_paragraph()
    tail.paragraph_format.space_before = Pt(0)
    tail.paragraph_format.space_after = Pt(0)
    tail.paragraph_format.line_spacing = Pt(1)
    tail.paragraph_format.page_break_before = False
    tail.paragraph_format.keep_with_next = False
    tail.add_run().font.size = Pt(1)
    snap = OxmlElement('w:snapToGrid')
    snap.set(qn('w:val'), '0')
    tail._p.get_or_add_pPr().append(snap)
    keep_item_analysis_connected(document)
    document.save(docx_path)
    audit_path = Path(docx_path).with_suffix('.images.json')
    audit_path.write_text(json.dumps({'manifest':str(manifest_path.resolve()), 'questions':mapping,
        'image_alignment':'left_top',
        'chinese_layout':'intro_with_first_item' if source_subject=='C' else None,
        'layout':{key:value for key,value in layout.items() if key!='sizes'},'material_layouts':material_layouts,
        'materials':sorted(inserted_groups)},ensure_ascii=False,indent=2),encoding='utf-8')
    return audit_path


def mark_missing_statistics(docx_path, chapter):
    document = Document(docx_path)
    if document.paragraphs:
        document.paragraphs[0].add_run('（待補施測統計）')
    paragraphs = list(document.paragraphs)
    seen = set()
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        text = paragraph.text
        if '本報告先分析' in text:
            paragraph.text = '本章整理題本、標準答案與評量向度。信度、有效樣本數及施測統計資料待補。'
        elif '整份試卷未作答' in text:
            paragraph.text = '有效樣本數與施測資料清理結果待提供，試題分析名詞定義如表4-1。'
        elif '全體學生通過率是待補' in text or '平均通過率是待補' in text:
            paragraph.text = '各評量向度與題號整理如下，施測後通過率待補；目前不做學力表現判讀。'
        elif '待補%' in text or text.strip()=='-':
            paragraph.text = text.replace('待補%', '待補') if text.strip()!='-' else '待補'
        for run in paragraph.runs:
            run.font.name = '微軟正黑體'
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微軟正黑體')
    document.save(docx_path)


def normalize_chapter6_layout(docx_path, subject):
    """長題號清單不能被固定列高遮住；矩陣同一向度的三列一起移頁。"""
    document = Document(docx_path)
    table = document.tables[0]
    seen = set()
    for row in table.rows:
        for cell in row.cells:
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            for paragraph in cell.paragraphs:
                # 範本的 18pt 文件格線會把 24pt 行高向上對齊成 36pt，擠出尾列。
                # 表格依自身行高排版，正文仍沿用原文件格線。
                ppr = paragraph._p.get_or_add_pPr()
                snap = ppr.find(qn('w:snapToGrid'))
                if snap is None:
                    snap = OxmlElement('w:snapToGrid')
                    ppr.append(snap)
                snap.set(qn('w:val'),'0')
    if subject in {'C','E'}:
        for row in table.rows:
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
    elif subject in {'M','S'}:
        seen_cells = set()
        for index,row in enumerate(table.rows):
            if index < 2:
                continue
            # 原範本首列約 40pt，複製到每一列會把 17 列矩陣擠成兩頁。
            # 合併格的長題號可在三列的總高度內換行。
            row.height = Pt(24)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            keep_next = (index-2)%3 < 2
            for cell in row.cells:
                if cell._tc in seen_cells:
                    continue
                seen_cells.add(cell._tc)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = keep_next
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
    document.save(docx_path)
