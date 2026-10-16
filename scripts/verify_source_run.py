"""獨立由原始 Excel 核對逐題統計與表 6-1，不用報告產生器的平均函式。"""
import argparse
from decimal import Decimal, ROUND_HALF_UP
import json
from pathlib import Path
import re

from docx import Document
import openpyxl
from report_images import item_tables
from qa_verify import verify_question_images


def mean_text(raw, numbers, group):
    value = sum(Decimal(str(raw[n][group])) for n in numbers)/Decimal(len(numbers))
    return str(value.quantize(Decimal('.01'), rounding=ROUND_HALF_UP))


def clean(value):
    return re.sub(r'\s+', '', str(value))


def unique_cells(row):
    result, seen = [], set()
    for cell in row.cells:
        if cell._tc not in seen:
            result.append(cell)
            seen.add(cell._tc)
    return result


def verify(output_dir, suffix='_AI產出'):
    output_dir = Path(output_dir)
    builds = json.loads((output_dir/f'產出清單{suffix}.json').read_text(encoding='utf-8'))
    result = {'passed':True, 'question_values_checked':0, 'domain_rate_cells_checked':0, 'reports':[], 'errors':[]}
    def require(condition, message):
        if not condition:
            result['errors'].append(message)
    for build in builds:
        cohort = build['cohort']
        image_checks = verify_question_images(build['chapter4'])
        require(not any(c[1]=='FAIL' for c in image_checks),f'{cohort} 圖片驗證失敗')
        if build['draft']:
            result['reports'].append({'cohort':cohort,'statistics':'pending','image_checks':image_checks})
            continue
        workbook = openpyxl.load_workbook(build['analysis_file'],data_only=True)
        sheet = workbook.worksheets[0]
        summary = {int(row[0]):row for row in list(sheet.values)[3:] if isinstance(row[0],(int,float))}
        raw_file = Path(build['analysis_file']).with_name(cohort+'_分析結果.xlsx')
        raw_book = openpyxl.load_workbook(raw_file,data_only=True)
        raw_sheet = raw_book['工作表1']
        raw = {n:{'hi':raw_sheet.cell(n+1,9).value,'all':raw_sheet.cell(n+1,4).value,
                  'lo':raw_sheet.cell(n+1,7).value} for n in summary}
        answers, domains = build['answers'], build['domains']
        table_map = item_tables(Document(build['chapter4']))
        grade_code = '6' if cohort.endswith('N') else re.sub(r'\D','',cohort)
        require(set(table_map)==set(summary),f'{cohort} 題數不符')
        for number, table in table_map.items():
            values = summary[number]
            require(table.cell(1,len(table.columns)-1).text.strip()==answers[number-1],f'{cohort} q{number} 答案錯誤')
            require(re.fullmatch(r'\d{3}-'+cohort[0]+grade_code+r'-[A-Z]\d-'+f'{number:02d}',table.cell(1,1).text) is not None,
                    f'{cohort} q{number} 試題代碼錯誤')
            for row_index, offset in [(6,5),(7,10),(8,15)]:
                cells = unique_cells(table.rows[row_index])
                actual = [c.text.strip() for c in cells[2:7]]
                expected = [f'{values[offset+i] or 0:.2f}' for i in range(5)]
                require(actual==expected,f'{cohort} q{number} 選答率列{row_index}錯誤: {actual} != {expected}')
            text = '\n'.join(c.text for c in unique_cells(table.rows[5]))
            rate = (Decimal(str(values[3]))*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP)
            require(f'通過率：{rate}%' in text,f'{cohort} q{number} 通過率錯誤')
            discrimination = f'{values[1]:.2f}'
            require(f'鑑別度：{discrimination}' in '\n'.join(c.text for c in unique_cells(table.rows[7])),f'{cohort} q{number} 鑑別度錯誤')
            require(str(raw_sheet.cell(number+1,2).value).upper()==answers[number-1],f'{cohort} 原始統計題序／答案錯誤 q{number}')
            result['question_values_checked'] += 1
        table = Document(build['chapter6']).tables[0]
        def compare(row, column, numbers, group):
            if not numbers:
                return
            actual, expected = table.cell(row,column).text.strip(), mean_text(raw,numbers,group)
            require(actual==expected,f'{cohort} 表6-1({row},{column}) {actual} != {expected}')
            result['domain_rate_cells_checked'] += 1
        if cohort[0]=='C':
            for row in range(2,len(table.rows)):
                label = clean(table.cell(row,0).text)
                numbers = list(summary) if label=='整體' else [n for n,d in enumerate(domains,1) if d[1]==label]
                require(bool(numbers),f'{cohort} 表6-1未知向度 {label}')
                for column,group in [(2,'hi'),(3,'all'),(4,'lo')]:compare(row,column,numbers,group)
        elif cohort[0]=='E':
            for row in range(2,len(table.rows)):
                main,sub = clean(table.cell(row,0).text),clean(table.cell(row,1).text)
                numbers = list(summary) if main=='整體' else [n for n,d in enumerate(domains,1)
                    if clean(d[0])==main and re.sub(r'^[AB]\d+','',clean(d[1]))==sub]
                require(bool(numbers),f'{cohort} 表6-1未知向度 {main}/{sub}')
                for column,group in [(3,'hi'),(4,'all'),(5,'lo')]:compare(row,column,numbers,group)
        else:
            cognitive = ['概念理解','程序執行','解題思考']
            for start_row in range(2,len(table.rows),3):
                label = clean(table.cell(start_row,0).text)
                content_numbers = list(summary) if label=='整體' else [n for n,d in enumerate(domains,1) if clean(d[0])==label]
                require(bool(content_numbers),f'{cohort} 表6-1未知向度 {label}')
                for row_offset,group in enumerate(['hi','all','lo']):
                    for index,name in enumerate(cognitive):
                        numbers = [n for n in content_numbers if domains[n-1][1]==name]
                        compare(start_row+row_offset,3+index*3,numbers,group)
                    compare(start_row+row_offset,12,content_numbers,group)
        workbook.close()
        raw_book.close()
        result['reports'].append({'cohort':cohort,'statistics':'checked','image_checks':image_checks})
    result['passed'] = not result['errors']
    return result


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--dir',required=True,type=Path)
    parser.add_argument('--output-suffix',default='_AI產出')
    args=parser.parse_args()
    result=verify(args.dir,args.output_suffix)
    (args.dir/f'獨立數值驗證{args.output_suffix}.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False,indent=2))
    if not result['passed']:raise SystemExit(1)
