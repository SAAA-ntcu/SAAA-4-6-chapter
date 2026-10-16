"""
SAAA-4-6-chapter: Unified Report Generation Engine
Automates Chapter 4 (肆、試題分析與教學建議) and Chapter 6 (陸、整體教學建議 / 表6-1 各評量向度通過率)
Supports 4 subjects: Chinese (C), Math (M), English (E), Science (S)
Supports standard cohorts (grades 3~8) and special cohorts (e.g., C7N, M7N, E7N, S7N).
"""

import os
import sys
import re
import copy
import argparse
from datetime import datetime
import pandas as pd
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Pt, Twips
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docxtpl import DocxTemplate

# Default Constants
CIRCLE_NUMS = {"1": "①", "2": "②", "3": "③", "4": "④"}
GRADE_CN_MAP = {1: "一", 2: "二", 3: "三", 4: "四", 5: "五", 6: "六", 7: "七", 8: "八", 9: "九"}
GRADE_CN_FULL = {
    3: "三年級", 4: "四年級", 5: "五年級", 6: "六年級", 7: "七年級", 8: "八年級"
}
OPTION_LABELS_36 = ["1", "2", "3", "4"]
OPTION_LABELS_78 = ["A", "B", "C", "D"]
OPT_MAP = {'A': 0, 'B': 1, 'C': 2, 'D': 3, '1': 0, '2': 1, '3': 2, '4': 3}
HEADING_CN = {1: "一", 2: "二", 3: "三", 4: "四", 5: "五", 6: "六", 7: "七", 8: "八", 9: "九", 10: "十"}
CN_NUMS = {1: "一", 2: "二", 3: "三", 4: "四", 5: "五", 6: "六", 7: "七", 8: "八", 9: "九", 10: "十", 11: "十一", 12: "十二"}

CATEGORY_A_NAMES_C = {"形音知識", "字詞知識", "語法知識", "修辭知識", "章法知識", "文體知識", "國學知識"}

SUBJECT_NAME_MAP = {
    "C": "國語",
    "M": "數學",
    "E": "英文",
    "S": "自然"
}

SUBJECT_FULL_NAMES = {
    "C": "國語文",
    "M": "數學",
    "E": "英語文",
    "S": "自然科學"
}

# ==============================================================================
# Formatting and Word XML Helpers
# ==============================================================================
def format_rate(val):
    if pd.isna(val) or val is None or val == "-":
        return "-"
    try:
        return str(int(round(float(val) * 100)))
    except (ValueError, TypeError):
        return str(val)

def format_disc(val):
    if pd.isna(val) or val is None or val == "-":
        return "0.00"
    try:
        return f"{float(val):.2f}"
    except (ValueError, TypeError):
        return str(val)

def format_avg_rate(val):
    if pd.isna(val) or val is None or val == "" or val == "-":
        return "-"
    try:
        return f"{float(val) + 1e-9:.2f}"
    except (ValueError, TypeError):
        return str(val)

def ans_display(ans_str, grade_val=7):
    if not ans_str or pd.isna(ans_str):
        return ""
    s = str(ans_str).strip()
    is_junior_high = False
    if isinstance(grade_val, str) and ("7" in grade_val or "8" in grade_val or "N" in grade_val):
        is_junior_high = True
    elif isinstance(grade_val, int) and grade_val >= 7:
        is_junior_high = True
    if is_junior_high:
        return s.upper()
    return CIRCLE_NUMS.get(s, s)

def set_cell_value(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, font_size_pt=12, template_cell=None, squeeze=False):
    p_list = cell.paragraphs
    for p in p_list[1:]:
        p._p.getparent().remove(p._p)
    p = p_list[0]
    
    t_pPr = None
    if template_cell is not None and len(template_cell.paragraphs) > 0:
        t_pPr = copy.deepcopy(template_cell.paragraphs[0]._p.get_or_add_pPr())
        
    p._p.clear()
    if t_pPr is not None:
        p._p.append(t_pPr)
    p.alignment = align
    
    if text is not None and text != "":
        run = p.add_run(str(text))
        run.font.bold = bold
        run.font.name = '微軟正黑體'
        run.font.size = Pt(font_size_pt)
        rPr = run._r.get_or_add_rPr()
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:eastAsia'), '微軟正黑體')
        rFonts.set(qn('w:ascii'), '微軟正黑體')
        rFonts.set(qn('w:hAnsi'), '微軟正黑體')
        rPr.append(rFonts)
        if squeeze:
            spacing = OxmlElement('w:spacing')
            spacing.set(qn('w:val'), '-20')
            rPr.append(spacing)

def set_paragraph_text_font(p, text):
    p.text = text
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.font.name = '微軟正黑體'
        rPr = run._r.get_or_add_rPr()
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:eastAsia'), '微軟正黑體')
        rFonts.set(qn('w:ascii'), '微軟正黑體')
        rFonts.set(qn('w:hAnsi'), '微軟正黑體')
        rPr.append(rFonts)

def set_cell_diagonal_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    tl2br = tcBorders.find(qn('w:tl2br'))
    if tl2br is None:
        tl2br = OxmlElement('w:tl2br')
        tl2br.set(qn('w:val'), 'single')
        tl2br.set(qn('w:sz'), '4')
        tl2br.set(qn('w:space'), '0')
        tl2br.set(qn('w:color'), 'BFBFBF')
        tcBorders.append(tl2br)

def set_cell_shading(cell, fill_hex="E0E0E0"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), fill_hex)
        tcPr.append(shd)
    else:
        shd.set(qn('w:fill'), fill_hex)

def add_cloned_row(table, src_tr):
    new_tr = copy.deepcopy(src_tr)
    for tc in new_tr.xpath('.//w:tc'):
        tcPr = tc.get_or_add_tcPr()
        vMerge = tcPr.find(qn('w:vMerge'))
        if vMerge is not None:
            tcPr.remove(vMerge)
        p_list = tc.xpath('.//w:p')
        for p in p_list[1:]:
            p.getparent().remove(p)
        p = p_list[0]
        pPr = p.find(qn('w:pPr'))
        p.clear()
        if pPr is not None:
            p.append(pPr)
    table._tbl.append(new_tr)
    return table.rows[-1]

def set_table_column_widths(table, col_widths):
    for row in table.rows:
        seen_cells = set()
        col_idx = 0
        for cell in row.cells:
            if cell in seen_cells:
                continue
            seen_cells.add(cell)
            tcPr = cell._tc.get_or_add_tcPr()
            gridSpan = tcPr.find(qn('w:gridSpan'))
            span = int(gridSpan.get(qn('w:val'))) if gridSpan is not None else 1
            cell_width = sum(col_widths[col_idx : col_idx + span])
            cell.width = cell_width
            col_idx += span

def update_table_grid_widths(table, col_widths):
    tblGrid = table._tbl.tblGrid
    if tblGrid is not None:
        tblGrid.clear()
    else:
        tblGrid = OxmlElement('w:tblGrid')
        table._tbl.insert(1, tblGrid)
    for w in col_widths:
        gridCol = OxmlElement('w:gridCol')
        twips_val = round(int(w) / 635)
        gridCol.set(qn('w:w'), str(twips_val))
        tblGrid.append(gridCol)

def set_table_cell_margins(table, top=60, bottom=60, left=45, right=45):
    tblPr = table._tbl.tblPr
    tblCellMar = tblPr.find(qn('w:tblCellMar'))
    if tblCellMar is None:
        tblCellMar = OxmlElement('w:tblCellMar')
        tblPr.append(tblCellMar)
    else:
        tblCellMar.clear()
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tblCellMar.append(node)

def format_total_questions(count):
    return f"共\n{count}\n題"

def format_domain_name_two_by_two(name):
    chunks = [name[i:i+2] for i in range(0, len(name), 2)]
    return '\n'.join(chunks)

def format_question_list(q_list, style="math"):
    if not q_list:
        return "", WD_ALIGN_PARAGRAPH.CENTER
    q_sorted = sorted(list(set(q_list)))
    if style == "science":
        text = '\n'.join(f'{q:02d}' for q in q_sorted)
        return text, WD_ALIGN_PARAGRAPH.CENTER
    if len(q_sorted) <= 2:
        text = '、'.join(f'{q:02d}' for q in q_sorted)
        return text, WD_ALIGN_PARAGRAPH.CENTER
    formatted_lines = []
    for idx in range(0, len(q_sorted), 2):
        pair = q_sorted[idx : idx + 2]
        line_text = '、'.join(f'{q:02d}' for q in pair)
        if idx + 2 < len(q_sorted):
            line_text += '、'
        formatted_lines.append(line_text)
    return '\n'.join(formatted_lines), WD_ALIGN_PARAGRAPH.LEFT

# Page break handlers for Chapter 4
def add_pagebreaks_in_tables_chinese(doc):
    body = doc.element.body
    first_q_table = None
    for child in body:
        if child.tag.endswith('}tbl'):
            first_q_table = child
            break

    if first_q_table is not None:
        for child in list(body):
            if child is first_q_table:
                break
            if child.tag.endswith('}p'):
                text = ''.join(t.text or '' for t in child.findall('.//' + qn('w:t')))
                has_pb = child.findall('.//' + qn('w:br'))
                has_page_br = any(br.get(qn('w:type')) == 'page' for br in has_pb)
                if not text.strip() and not has_page_br:
                    body.remove(child)

    heading_pattern = re.compile(r'^\(.\).+')
    prev_was_heading = False
    for child in list(body):
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if tag == 'p':
            text = ''.join(t.text or '' for t in child.findall('.//' + qn('w:t')))
            if heading_pattern.match(text.strip()):
                prev_was_heading = True
            elif text.strip():
                prev_was_heading = False
        elif tag == 'tbl':
            is_first_table = (child is first_q_table)
            if not is_first_table and not prev_was_heading:
                first_row = child.find(qn('w:tr'))
                if first_row is not None:
                    first_cell = first_row.find(qn('w:tc'))
                    if first_cell is not None:
                        first_para = first_cell.find(qn('w:p'))
                        if first_para is not None:
                            pPr = first_para.find(qn('w:pPr'))
                            if pPr is None:
                                pPr = OxmlElement('w:pPr')
                                first_para.insert(0, pPr)
                            pPr.append(OxmlElement('w:pageBreakBefore'))
            prev_was_heading = False

    for child in list(body):
        if child.tag.endswith('}p'):
            text = ''.join(t.text or '' for t in child.findall('.//' + qn('w:t')))
            has_pb = child.findall('.//' + qn('w:br'))
            has_page_br = any(br.get(qn('w:type')) == 'page' for br in has_pb)
            if not text.strip() and not has_page_br:
                body.remove(child)

    subheading_pattern = re.compile(r'^\([二三四五六七八九十]\).+')
    major_heading_pattern = re.compile(r'^二、')
    for p in doc.paragraphs:
        text = p.text.strip()
        if subheading_pattern.match(text) or major_heading_pattern.match(text):
            pPr = p.paragraph_format.element.get_or_add_pPr()
            if pPr.find(qn('w:pageBreakBefore')) is None:
                pPr.append(OxmlElement('w:pageBreakBefore'))

def add_pagebreaks_in_tables_continuous(doc):
    body = doc.element.body
    tables = body.findall(qn('w:tbl'))
    first_q_table = tables[1] if len(tables) > 1 else None
    if first_q_table is not None:
        for child in list(body):
            if child is first_q_table:
                break
            if child.tag.endswith('}p'):
                text = ''.join(t.text or '' for t in child.findall('.//' + qn('w:t')))
                if not text.strip():
                    body.remove(child)

    for tbl in tables[1:]:
        first_row = tbl.find(qn('w:tr'))
        if first_row is None: continue
        first_cell = first_row.find(qn('w:tc'))
        if first_cell is None: continue
        first_para = first_cell.find(qn('w:p'))
        if first_para is None: continue
        pPr = first_para.find(qn('w:pPr'))
        if pPr is None:
            pPr = OxmlElement('w:pPr')
            first_para.insert(0, pPr)
        pPr.append(OxmlElement('w:pageBreakBefore'))

    for child in list(body):
        if child.tag.endswith('}p'):
            text = ''.join(t.text or '' for t in child.findall('.//' + qn('w:t')))
            if not text.strip():
                body.remove(child)

# ==============================================================================
# Excel Readers for Item Analysis and Dimensions
# ==============================================================================
def read_analysis_file(path, grade_target=None):
    xls = pd.ExcelFile(path)
    results = {}
    for sheet_name in xls.sheet_names:
        sheet_clean = str(sheet_name).strip()
        g_key = sheet_clean
        m = re.search(r'\d+', sheet_clean)
        if m:
            g_key = m.group(0)
            if "N" in sheet_clean.upper():
                g_key += "N"
        
        df = pd.read_excel(xls, sheet_name=sheet_name, header=None)
        if len(df) < 4:
            continue
        meta_row = df.iloc[0]
        title = str(meta_row.iloc[0]) if pd.notna(meta_row.iloc[0]) else ""
        
        sample_str = ""
        for cell_val in meta_row:
            if pd.notna(cell_val) and "有效樣本數" in str(cell_val):
                sample_str = str(cell_val)
                break
        sample_m = re.search(r"有效樣本數[：:](.+)", sample_str)
        valid_sample = sample_m.group(1).strip() if sample_m else ""

        alpha_str = ""
        for cell_val in meta_row:
            if pd.notna(cell_val) and ("Cronbach" in str(cell_val) or "α" in str(cell_val)):
                alpha_str = str(cell_val)
                break
        alpha_m = re.search(r"Cronbach.s\s*α?[：:]?\s*(.+)", alpha_str)
        cronbach = alpha_m.group(1).strip() if alpha_m else "-"

        questions = []
        for idx in range(3, len(df)):
            row = df.iloc[idx]
            q_num = row.iloc[0]
            if pd.isna(q_num): continue
            try:
                q_num = int(float(q_num))
            except (ValueError, TypeError):
                continue
            disc = row.iloc[1]
            diff = row.iloc[2]
            pass_rate = row.iloc[3]
            correct_ans = str(row.iloc[4]).strip() if pd.notna(row.iloc[4]) else None
            
            is_invalid = False
            if correct_ans is None or correct_ans.lower() == "nan" or "不予計分" in correct_ans:
                is_invalid = True
                correct_ans = "不予計分"

            all_rates = [float(row.iloc[c]) if pd.notna(row.iloc[c]) else 0.0 for c in range(5, 10)]
            hi_rates = [float(row.iloc[c]) if pd.notna(row.iloc[c]) else 0.0 for c in range(10, 15)]
            lo_rates = [float(row.iloc[c]) if pd.notna(row.iloc[c]) else 0.0 for c in range(15, 20)]

            questions.append({
                "num": q_num,
                "disc": disc,
                "diff": diff,
                "pass_rate": pass_rate,
                "correct_ans": correct_ans,
                "all_rates": all_rates,
                "hi_rates": hi_rates,
                "lo_rates": lo_rates,
                "invalid": is_invalid
            })
        
        # Check if an unrounded analysis result file exists alongside (e.g. C7N_分析結果.xlsx)
        base_dir = os.path.dirname(path)
        candidates = [
            path.replace("試題分析總表", "分析結果"),
            path.replace("_試題分析總表", "_分析結果"),
            os.path.join(base_dir, f"{sheet_clean}_分析結果.xlsx"),
            os.path.join(base_dir, f"{g_key}_分析結果.xlsx"),
        ]
        res_file = None
        for cand in candidates:
            if cand != path and os.path.exists(cand):
                res_file = cand
                break

        if res_file:
            try:
                import openpyxl
                wb_res = openpyxl.load_workbook(res_file, data_only=True)
                target_sheet = '工作表1' if '工作表1' in wb_res.sheetnames else wb_res.sheetnames[0]
                ws_res = wb_res[target_sheet]
                q_res_map = {}
                for r in range(2, ws_res.max_row + 1):
                    q_idx = r - 1
                    try:
                        p_val = ws_res.cell(r, 4).value
                        lo_val = ws_res.cell(r, 7).value
                        hi_val = ws_res.cell(r, 9).value
                        if hi_val is not None and lo_val is not None:
                            q_res_map[q_idx] = {
                                'hi': float(hi_val),
                                'lo': float(lo_val),
                                'all': float(p_val) if p_val is not None else None
                            }
                    except (ValueError, TypeError):
                        pass
                for q in questions:
                    if q["num"] in q_res_map:
                        q["raw_hi_pass"] = q_res_map[q["num"]]['hi']
                        q["raw_lo_pass"] = q_res_map[q["num"]]['lo']
                        if q_res_map[q["num"]]['all'] is not None:
                            q["raw_pass_rate"] = q_res_map[q["num"]]['all']
            except Exception as e:
                print(f"警告: 讀取未四捨五入分析結果失敗 ({res_file}): {e}")

        parsed = {"title": title, "valid_sample": valid_sample, "cronbach": cronbach, "questions": questions}
        results[sheet_clean] = parsed
        results[g_key] = parsed
        try:
            results[int(g_key)] = parsed
        except (ValueError, TypeError):
            pass

    return results

def read_answers_and_domains(path, subject, grade_target=None):
    xls = pd.ExcelFile(path)
    ans_sheet = None
    for s in xls.sheet_names:
        if "答案" in s:
            ans_sheet = s
            break
    if ans_sheet is None:
        ans_sheet = xls.sheet_names[0]

    dim_sheet = None
    for s in xls.sheet_names:
        if "向度" in s or "指標" in s:
            dim_sheet = s
            break
    if dim_sheet is None:
        dim_sheet = xls.sheet_names[-1]

    df_ans = pd.read_excel(xls, sheet_name=ans_sheet, header=0)
    df_dim = pd.read_excel(xls, sheet_name=dim_sheet, header=0)

    # Detect grades present in answer sheet
    grades = {}
    subject_code = subject.upper()[0]

    for col in df_ans.columns:
        c_str = str(col).strip()
        # Look for e.g. C3, C7N, M4, E7, S5
        m = re.search(r'[A-Za-z]?(\d+)(N?)', c_str)
        if m:
            g_num = m.group(1)
            suffix = m.group(2)
            g_key = f"{g_num}{suffix}"
            if g_key not in grades:
                grades[g_key] = {"ans_col": col}
        elif "答案" in c_str or "ans" in c_str.lower():
            grades["single"] = {"ans_col": col}

    # If single cohort and no grade regex found
    if not grades:
        grades["single"] = {"ans_col": df_ans.columns[-1]}

    # Map dimensions for each detected grade
    for g_key, g_info in list(grades.items()):
        ans_col = g_info["ans_col"]
        ans_list = [str(v).strip() if pd.notna(v) else "" for v in df_ans[ans_col].tolist()]
        g_info["answers"] = ans_list

        # Detect dimension columns in df_dim
        dim_data = []
        if subject_code == "C":
            # Chinese has domain names like 'A11 形音知識'
            col_target = None
            for c in df_dim.columns:
                c_str = str(c)
                if g_key != "single" and g_key in c_str:
                    col_target = c
                    break
                elif g_key != "single" and GRADE_CN_MAP.get(int(re.sub(r'\D', '', g_key) or 0), "") in c_str:
                    col_target = c
                    break
            if col_target is None:
                # Pick any column with 年級 or 向度 or 指標
                cand = [c for c in df_dim.columns if '年級' in str(c) or '向度' in str(c) or '指標' in str(c)]
                col_target = cand[0] if cand else df_dim.columns[-1]

            for i in range(len(df_dim)):
                raw = str(df_dim[col_target].iloc[i]).strip() if pd.notna(df_dim[col_target].iloc[i]) else ""
                code, name = raw.split(" ", 1) if " " in raw else ("", raw)
                dim_data.append((code, name))
            g_info["domains"] = dim_data

        elif subject_code in ["M", "E", "S"]:
            d1_col, d2_col = None, None
            for c in df_dim.columns:
                c_str = str(c)
                if (g_key != "single" and g_key in c_str) or ("1" in c_str and "向度" in c_str):
                    if "1" in c_str and ("向度" in c_str or "指標" in c_str):
                        d1_col = c
                    elif "2" in c_str and ("向度" in c_str or "指標" in c_str):
                        d2_col = c
            if d1_col is None or d2_col is None:
                dim_cols = [c for c in df_dim.columns if '向度' in str(c) or '指標' in str(c)]
                if len(dim_cols) >= 2:
                    d1_col, d2_col = dim_cols[0], dim_cols[1]
                elif len(dim_cols) == 1:
                    d1_col, d2_col = dim_cols[0], dim_cols[0]
                else:
                    d1_col, d2_col = df_dim.columns[0], df_dim.columns[1] if len(df_dim.columns) > 1 else df_dim.columns[0]

            for i in range(len(df_dim)):
                v1 = str(df_dim[d1_col].iloc[i]).strip() if pd.notna(df_dim[d1_col].iloc[i]) else ""
                v2 = str(df_dim[d2_col].iloc[i]).strip() if pd.notna(df_dim[d2_col].iloc[i]) else ""
                dim_data.append((v1, v2))
            g_info["domains"] = dim_data

    return grades

# ==============================================================================
# Chapter 4 Question Dict Builder
# ==============================================================================
def build_question_dict(q_data, code, num, c_ans, domain, cognitive="", opt_labels=OPTION_LABELS_78, is_invalid=False):
    options = opt_labels
    correct = str(c_ans).strip().upper() if pd.notna(c_ans) else ""
    correct_parts = [x.strip().upper() for x in correct.split('、')] if correct else []
    
    h = {}
    for i, opt in enumerate(options[:4]):
        h[f"h{i}"] = f"{opt}*" if (opt.upper() in correct_parts and not is_invalid) else opt
    if q_data.get('missing_statistics'):
        context = {'number':num, 'code':code, 'domain':domain, 'cognitive':cognitive,
            'sub_domain':cognitive, 'ans':correct, 'correct_ans':correct,
            **{f'h{i+1}':h[f'h{i}'] for i in range(4)}}
        keys = ['discrimination','disc','pass_rate','hi_pass','lo_pass']
        keys += [f'{g}_{n}' for g in ['all','hi','lo'] for n in [1,2,3,4,9]]
        keys += [f'w{i}_{k}' for i in range(1,4) for k in ['ans','opt','rate']]
        context.update({k:'待補' for k in keys})
        return context
        
    disc_val = float(q_data["disc"]) if pd.notna(q_data["disc"]) else 0.0
    pass_val = float(q_data["pass_rate"]) if pd.notna(q_data["pass_rate"]) else 0.0
    
    all_rates = q_data["all_rates"]
    hi_rates = q_data["hi_rates"]
    lo_rates = q_data["lo_rates"]
    
    correct_indices = [i for i, opt in enumerate(options[:4]) if opt.upper() in correct_parts]
    hi_pass = sum(hi_rates[i] for i in correct_indices) if correct_indices else 0.0
    lo_pass = sum(lo_rates[i] for i in correct_indices) if correct_indices else 0.0
    
    wrong_indices = [i for i in range(len(options[:4])) if i not in correct_indices]
    wrong_sorted = sorted(wrong_indices)
    
    disc_fmt = format_disc(disc_val) if not is_invalid else "-"
    pass_fmt = format_rate(pass_val) if not is_invalid else "-"
    ans_fmt = correct if not is_invalid else "不予計分"
    
    q_ctx = {
        "number": num,
        "code": code,
        "domain": domain,
        "cognitive": cognitive,
        "sub_domain": cognitive,
        "ans": ans_fmt,
        "correct_ans": ans_fmt,
        "discrimination": disc_fmt,
        "disc": disc_fmt,
        "pass_rate": pass_fmt,
        "hi_pass": format_rate(hi_pass) if not is_invalid else "-",
        "lo_pass": format_rate(lo_pass) if not is_invalid else "-",
        "h1": h.get("h0", "A"),
        "h2": h.get("h1", "B"),
        "h3": h.get("h2", "C"),
        "h4": h.get("h3", "D"),
        "all_1": f"{all_rates[0]:.2f}" if not is_invalid else "-",
        "all_2": f"{all_rates[1]:.2f}" if not is_invalid else "-",
        "all_3": f"{all_rates[2]:.2f}" if not is_invalid else "-",
        "all_4": f"{all_rates[3]:.2f}" if not is_invalid else "-",
        "all_9": f"{all_rates[4]:.2f}" if not is_invalid else "-",
        "hi_1": f"{hi_rates[0]:.2f}" if not is_invalid else "-",
        "hi_2": f"{hi_rates[1]:.2f}" if not is_invalid else "-",
        "hi_3": f"{hi_rates[2]:.2f}" if not is_invalid else "-",
        "hi_4": f"{hi_rates[3]:.2f}" if not is_invalid else "-",
        "hi_9": f"{hi_rates[4]:.2f}" if not is_invalid else "-",
        "lo_1": f"{lo_rates[0]:.2f}" if not is_invalid else "-",
        "lo_2": f"{lo_rates[1]:.2f}" if not is_invalid else "-",
        "lo_3": f"{lo_rates[2]:.2f}" if not is_invalid else "-",
        "lo_4": f"{lo_rates[3]:.2f}" if not is_invalid else "-",
        "lo_9": f"{lo_rates[4]:.2f}" if not is_invalid else "-",
    }
    
    for wi in range(3):
        if wi < len(wrong_sorted) and not is_invalid:
            w_idx = wrong_sorted[wi]
            w_opt = options[w_idx]
            q_ctx[f"w{wi + 1}_ans"] = w_opt
            q_ctx[f"w{wi + 1}_opt"] = w_opt
            q_ctx[f"w{wi + 1}_rate"] = str(int(round(all_rates[w_idx] * 100)))
        else:
            q_ctx[f"w{wi + 1}_ans"] = "-"
            q_ctx[f"w{wi + 1}_opt"] = "-"
            q_ctx[f"w{wi + 1}_rate"] = "-"
            
    return q_ctx

# ==============================================================================
# SUBJECT DISPATCHERS: CHAPTER 4
# ==============================================================================
def generate_chapter_4(subject, grade, year, answers, domains, analysis, template_path, output_path, images_manifest=None):
    subject_code = subject.upper()[0]
    grade_str = str(grade)
    is_entrant = ("N" in grade_str.upper()) or ("新生" in grade_str)
    grade_num = int(re.sub(r'\D', '', grade_str) or 7)
    grade_cn = GRADE_CN_MAP.get(grade_num, "七")
    opt_labels = OPTION_LABELS_78 if grade_num >= 7 else OPTION_LABELS_36
    
    # Internal question code uses '6' for 7N entrant testing 6th grade curriculum
    grade_code = "6" if is_entrant else str(grade_num)

    subj_full = SUBJECT_FULL_NAMES.get(subject_code, SUBJECT_NAME_MAP.get(subject_code, subject))
    if is_entrant:
        exam_name = f"{year}年桃園市{subj_full}七年級新生學習能力檢測"
    elif grade_num >= 7:
        exam_name = f"{year}年國民中學{subj_full}{grade_cn}年級學生學習能力檢測"
    else:
        exam_name = f"{year}年國民小學{subj_full}{grade_cn}年級學生學習能力檢測"

    print(f"[{subject_code}] 產生第四章: 年級={grade_str}, 試題數={len(analysis['questions'])}, 樣板={os.path.basename(template_path)}")

    if subject_code == "C":
        questions_ctx = []
        for q_data in analysis["questions"]:
            num = q_data["num"]
            d_code, d_name = domains[num - 1] if num - 1 < len(domains) else ("", "")
            c_ans = answers[num - 1] if num - 1 < len(answers) else ""
            code = f"{year}-C{grade_code}-{d_code}1-{str(num).zfill(2)}"
            q_ctx = build_question_dict(q_data, code, num, c_ans, d_name, cognitive="", opt_labels=opt_labels, is_invalid=q_data["invalid"])
            q_ctx["domain_code"] = d_code
            q_ctx["show_heading"] = False
            q_ctx["show_category_b"] = False
            q_ctx["heading_text"] = ""
            questions_ctx.append(q_ctx)

        questions_ctx.sort(key=lambda x: (x.get("domain_code", ""), x["number"]))

        current_cat = None
        cat_counter = 0
        seen_domains = set()
        seen_category_b = False
        for q_ctx in questions_ctx:
            domain = q_ctx["domain"]
            cat = "A" if domain in CATEGORY_A_NAMES_C else "B"
            if cat != current_cat:
                current_cat = cat
                cat_counter = 0
                seen_domains = set()
            if cat == "B" and not seen_category_b:
                q_ctx["show_category_b"] = True
                seen_category_b = True
            if domain not in seen_domains:
                seen_domains.add(domain)
                cat_counter += 1
                q_ctx["show_heading"] = True
                q_ctx["heading_text"] = f"({HEADING_CN.get(cat_counter, str(cat_counter))}){domain}"
            else:
                q_ctx["show_heading"] = False
                q_ctx["heading_text"] = ""

        seen_a, seen_b = [], []
        legend_parts = []
        seen = set()
        for c, n in domains:
            if n not in seen and n:
                seen.add(n)
                legend_parts.append(f"{c}：{n}" if c else n)
            if n in CATEGORY_A_NAMES_C:
                if n not in seen_a and n: seen_a.append(n)
            else:
                if n not in seen_b and n: seen_b.append(n)

        context = {
            "year": str(year),
            "subject": "國語",
            "grade_cn": grade_cn,
            "grade": grade_str,
            "cronbach_s_alpha": analysis["cronbach"],
            "valid_sample": analysis["valid_sample"],
            "exam_name": exam_name,
            "domain_legend": "；".join(legend_parts),
            "domain_intro": "本試題卷評量學生語文知識及閱讀理解兩項知能：",
            "category_a_title": "一、語文知識：",
            "category_a_list": f"{grade_cn}年級評量的語文知識計有：{'、'.join(seen_a)}等{CN_NUMS.get(len(seen_a), str(len(seen_a)))}項。",
            "category_b_title": "二、閱讀理解：",
            "category_b_list": f"{grade_cn}年級評量的閱讀理解計有：{'、'.join(seen_b)}等{CN_NUMS.get(len(seen_b), str(len(seen_b)))}項。",
            "questions": questions_ctx,
        }

        doc = DocxTemplate(template_path)
        doc.render(context)

        # Update preamble paragraph 1 text
        if len(doc.paragraphs) > 1:
            p1 = doc.paragraphs[1]
            if "本報告先分析" in p1.text:
                p1_text = f"本報告先分析「{exam_name}」結果，再針對各試題提供相關建議，本試卷信度Cronbach's alpha值為{analysis['cronbach']}，具一致性、穩定性及可靠性。"
                set_paragraph_text_font(p1, p1_text)

        # Chinese paragraph indentation and breaks
        body = doc.element.body
        for child in list(body):
            if child.tag.endswith('}p'):
                text = ''.join(t.text or '' for t in child.findall('.//' + qn('w:t')))
                if '評量的' in text:
                    pPr = child.find(qn('w:pPr'))
                    if pPr is None:
                        pPr = OxmlElement('w:pPr')
                        child.insert(0, pPr)
                    ind = OxmlElement('w:ind')
                    ind.set(qn('w:leftChars'), '350')
                    ind.set(qn('w:left'), '840')
                    ind.set(qn('w:firstLineChars'), '200')
                    ind.set(qn('w:firstLine'), '480')
                    pPr.append(ind)
                    break

        for child in list(body):
            if child.tag.endswith('}p'):
                text = ''.join(t.text or '' for t in child.findall('.//' + qn('w:t')))
                if '本試題卷評量學生' in text:
                    pb_para = OxmlElement('w:p')
                    pb_run = OxmlElement('w:r')
                    pb_br = OxmlElement('w:br')
                    pb_br.set(qn('w:type'), 'page')
                    pb_run.append(pb_br)
                    pb_para.append(pb_run)
                    child.addprevious(pb_para)
                    break

        add_pagebreaks_in_tables_chinese(doc)
        doc.save(output_path)

    elif subject_code == "M":
        domain_map = {
            "數與計算": "A", "量與實測": "B", "空間與形狀": "C", "關係": "D", "資料與不確定性": "E",
            "數與量": "A", "空間與形狀與坐標幾何": "B", "代數與函數": "C"
        }
        cog_map = {"概念理解": "1", "程序執行": "2", "解題思考": "3"}

        questions_ctx = []
        for q_data in analysis["questions"]:
            num = q_data["num"]
            d1, d2 = domains[num - 1] if num - 1 < len(domains) else ("", "")
            c_ans = answers[num - 1] if num - 1 < len(answers) else ""
            c1 = domain_map.get(d1, "A")
            c2 = cog_map.get(d2, "1")
            code = f"{year}-M{grade_code}-{c1}{c2}-{str(num).zfill(2)}"
            q_ctx = build_question_dict(q_data, code, num, c_ans, d1, cognitive=d2, opt_labels=opt_labels, is_invalid=q_data["invalid"])
            questions_ctx.append(q_ctx)

        seen = []
        for d1, d2 in domains:
            if d1 and d1 not in seen: seen.append(d1)
        code_map = sorted(((domain_map.get(d, '?'), d) for d in seen), key=lambda x: x[0])
        legend = "；".join(f"{code}：{d}" for code, d in code_map)

        context = {
            "year": str(year),
            "subject": "數學",
            "grade_cn": grade_cn,
            "grade": grade_str,
            "exam_name": exam_name,
            "cronbach_s_alpha": analysis["cronbach"],
            "valid_sample": analysis["valid_sample"],
            "domain_legend": legend,
            "questions": questions_ctx,
        }

        doc = DocxTemplate(template_path)
        doc.render(context)

        # Update preamble paragraph 1 text
        if len(doc.paragraphs) > 1:
            p1 = doc.paragraphs[1]
            if "本報告先分析" in p1.text:
                p1_text = f"本報告先分析「{exam_name}」結果，再針對各試題提供相關建議，本試卷信度Cronbach's alpha值為{analysis['cronbach']}，具一致性、穩定性及可靠性。"
                set_paragraph_text_font(p1, p1_text)

        add_pagebreaks_in_tables_continuous(doc)
        doc.save(output_path)

    elif subject_code == "E":
        parsed_domains = []
        for d1, d2 in domains:
            m = re.match(r'^([AB]\d+)(.+)$', str(d2).strip())
            c, n = (m.group(1), m.group(2)) if m else ("", str(d2).strip())
            parsed_domains.append((str(d1).strip(), c, n))

        questions_ctx = []
        for q_data in analysis["questions"]:
            num = q_data["num"]
            d1, sub_code, sub_name = parsed_domains[num - 1] if num - 1 < len(parsed_domains) else ("", "", "")
            c_ans = answers[num - 1] if num - 1 < len(answers) else ""
            code = f"{year}-E{grade_code}-{sub_code}-{str(num).zfill(2)}"
            q_ctx = build_question_dict(q_data, code, num, c_ans, d1, cognitive=sub_name, opt_labels=opt_labels, is_invalid=q_data["invalid"])
            q_ctx["sub_code"] = sub_code
            questions_ctx.append(q_ctx)

        seen_subs = {}
        for d1, sub_code, sub_name in parsed_domains:
            if d1 not in seen_subs: seen_subs[d1] = []
            if sub_code and (sub_code, sub_name) not in seen_subs[d1]:
                seen_subs[d1].append((sub_code, sub_name))

        parts = []
        main_codes = {"聽力": "A", "閱讀": "B"}
        for d1 in ["聽力", "閱讀"]:
            if d1 in seen_subs:
                sub_str = "，".join(f"{sc}：{sn}" for sc, sn in seen_subs[d1])
                parts.append(f"{main_codes.get(d1, d1)}：{d1}（{sub_str}）")
        legend = "；".join(parts)

        context = {
            "year": str(year),
            "subject": "英文",
            "grade_cn": grade_cn,
            "grade": grade_str,
            "exam_name": exam_name,
            "cronbach_s_alpha": analysis["cronbach"],
            "valid_sample": analysis["valid_sample"],
            "domain_legend": legend,
            "questions": questions_ctx,
        }

        doc = DocxTemplate(template_path)
        doc.render(context)

        # Update preamble paragraph 1 text
        if len(doc.paragraphs) > 1:
            p1 = doc.paragraphs[1]
            if "本報告先分析" in p1.text:
                p1_text = f"本報告先分析「{exam_name}」結果，再針對各試題提供相關建議，本試卷信度Cronbach's alpha值為{analysis['cronbach']}，具一致性、穩定性及可靠性。"
                set_paragraph_text_font(p1, p1_text)

        add_pagebreaks_in_tables_continuous(doc)
        doc.save(output_path)

    elif subject_code == "S":
        seen_codes = set()
        for d1, _ in domains:
            if d1:
                m = re.search(r"[（(](.+?)[）)]$", str(d1))
                if m: seen_codes.add(m.group(1))
        sorted_codes = sorted(list(seen_codes))
        domain_mapping = {c: chr(ord('A') + i) for i, c in enumerate(sorted_codes)}

        def parse_sci_d2(raw):
            if not raw: return "?", ""
            m = re.match(r"([A-D])\s*(.+)", str(raw))
            if m:
                mapping = {"A": "1", "B": "2", "C": "3", "D": "4"}
                return mapping.get(m.group(1).upper(), m.group(1)), m.group(2)
            return "?", str(raw)

        questions_ctx = []
        for q_data in analysis["questions"]:
            num = q_data["num"]
            d1_raw, d2_raw = domains[num - 1] if num - 1 < len(domains) else ("", "")
            c_ans = answers[num - 1] if num - 1 < len(answers) else ""
            m = re.search(r"[（(](.+?)[）)]$", str(d1_raw))
            orig_c = m.group(1) if m else "?"
            d_code = domain_mapping.get(orig_c, orig_c)
            d_name = re.sub(r"[（(].+?[）)]$", "", str(d1_raw)).strip()
            cog_code, cog_name = parse_sci_d2(d2_raw)

            code = f"{year}-S{grade_code}-{d_code}{cog_code}-{str(num).zfill(2)}"
            q_ctx = build_question_dict(q_data, code, num, c_ans, d_name, cognitive=cog_name, opt_labels=opt_labels, is_invalid=q_data["invalid"])
            q_ctx["domain_code"] = d_code
            questions_ctx.append(q_ctx)

        seen_d1 = []
        for d1, _ in domains:
            if d1 and d1 not in seen_d1: seen_d1.append(d1)
        legend_items = []
        for name in seen_d1:
            m = re.search(r"[（(](.+?)[）)]$", str(name))
            orig_c = m.group(1) if m else "?"
            mc = domain_mapping.get(orig_c, orig_c)
            dn = re.sub(r"[（(].+?[）)]$", "", str(name)).strip()
            legend_items.append((mc, dn))
        legend_items.sort(key=lambda x: x[0])
        domain_legend = "；".join(f"{mc}：{dn}" for mc, dn in legend_items)

        context = {
            "year": str(year),
            "subject": "自然",
            "grade_cn": grade_cn,
            "grade": grade_str,
            "exam_name": exam_name,
            "cronbach_s_alpha": analysis["cronbach"],
            "valid_sample": analysis["valid_sample"],
            "domain_legend": domain_legend,
            "questions": questions_ctx,
        }

        doc = DocxTemplate(template_path)
        doc.render(context)

        # Update preamble paragraph 1 text
        if len(doc.paragraphs) > 1:
            p1 = doc.paragraphs[1]
            if "本報告先分析" in p1.text:
                p1_text = f"本報告先分析「{exam_name}」結果，再針對各試題提供相關建議，本試卷信度Cronbach's alpha值為{analysis['cronbach']}，具一致性、穩定性及可靠性。"
                set_paragraph_text_font(p1, p1_text)

        add_pagebreaks_in_tables_continuous(doc)
        doc.save(output_path)

    if analysis.get('draft'):
        from report_images import mark_missing_statistics
        mark_missing_statistics(output_path, 4)
    if images_manifest:
        from report_images import attach_question_images
        attach_question_images(output_path, images_manifest)
    else:
        from report_images import keep_item_analysis_connected
        document=docx.Document(output_path)
        keep_item_analysis_connected(document)
        document.save(output_path)
    from proper_names import format_counties
    format_counties(output_path)
    print(f" -> 產出成功: {output_path}")

# ==============================================================================
# SUBJECT DISPATCHERS: CHAPTER 6 (TABLE 6-1)
# ==============================================================================
def generate_chapter_6(subject, grade, year, answers, domains, analysis, template_path, output_path):
    subject_code = subject.upper()[0]
    grade_str = str(grade)
    is_entrant = ("N" in grade_str.upper()) or ("新生" in grade_str)
    grade_num = int(re.sub(r'\D', '', grade_str) or 7)
    grade_cn = GRADE_CN_MAP.get(grade_num, "七")

    print(f"[{subject_code}] 產生第六章: 年級={grade_str}, 試題數={len(analysis['questions'])}, 樣板={os.path.basename(template_path)}")

    # Pre-calculate pass rates dictionary
    questions_data = {}
    for q in analysis["questions"]:
        num = q["num"]
        c_ans = answers[num - 1] if num - 1 < len(answers) else ""
        c_parts = [x.strip().upper() for x in str(c_ans).split('、')] if c_ans else []
        c_indices = [OPT_MAP[x] for x in c_parts if x in OPT_MAP]
        
        hi_pass = q.get("raw_hi_pass")
        if hi_pass is None:
            hi_pass = sum(q["hi_rates"][i] for i in c_indices) if c_indices else 0.0
            
        lo_pass = q.get("raw_lo_pass")
        if lo_pass is None:
            lo_pass = sum(q["lo_rates"][i] for i in c_indices) if c_indices else 0.0
            
        pass_val = q.get("raw_pass_rate")
        if pass_val is None:
            pass_val = float(q["pass_rate"]) if pd.notna(q["pass_rate"]) else (float('nan') if analysis.get('draft') else 0.0)
        
        questions_data[num] = {
            'hi': hi_pass,
            'all': pass_val,
            'lo': lo_pass
        }

    overall_pass_val = sum(q['all'] for q in questions_data.values()) / max(len(questions_data), 1)
    overall_pass_pct = '待補' if analysis.get('draft') else int(round(overall_pass_val * 100))

    if subject_code == "C":
        # Group domains
        domain_questions = {}
        for num, (d_code, d_name) in enumerate(domains, start=1):
            if not d_name: continue
            if d_name not in domain_questions:
                domain_questions[d_name] = {'prefix': d_code, 'questions': []}
            domain_questions[d_name]['questions'].append(num)

        sorted_domains = sorted(domain_questions.items(), key=lambda x: x[1]['prefix'])

        doc = docx.Document(template_path)
        # Update paragraph 1 and 2
        domain_names = [d[0] for d in sorted_domains]
        list_str = '」、「'.join(domain_names[:-1]) + '」與「' + domain_names[-1] if len(domain_names) > 1 else (domain_names[0] if domain_names else "")
        if is_entrant:
            p1_text = f"國語文七年級新生學力檢測，以「選擇題型」的方式，檢測「{list_str}」等{len(domain_names)}項能力。"
            p2_text = f"七年級新生國語文學力檢測全體學生通過率是{overall_pass_pct}%，各向度施測結果統計如下表6-1所示："
        else:
            p1_text = f"國語文{grade_cn}年級學力檢測，以「選擇題型」的方式，檢測「{list_str}」等{len(domain_names)}項能力。"
            p2_text = f"國語文{grade_cn}年級學力檢測全體學生通過率是{overall_pass_pct}%，各向度施測結果統計如下表6-1所示："
        if len(doc.paragraphs) > 1:
            set_paragraph_text_font(doc.paragraphs[1], p1_text)
        if len(doc.paragraphs) > 2:
            set_paragraph_text_font(doc.paragraphs[2], p2_text)

        for p in doc.paragraphs:
            if re.match(r'^表\s*6-1', p.text.strip()):
                set_paragraph_text_font(p, "表6-1　整體學生及高分組、低分組學生於各評量向度之通過率")
                break

        table = doc.tables[0]
        num_domains = len(sorted_domains)
        diff = 11 - num_domains
        if diff > 0 and len(table.rows) >= 14:
            for _ in range(diff):
                tr = table.rows[12]._tr
                table._tbl.remove(tr)

        # Style headers
        for r_idx in [0, 1]:
            for cell in table.rows[r_idx].cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                    p.paragraph_format.line_spacing = Pt(24)
                    for run in p.runs:
                        run.font.name = '微軟正黑體'
                        rPr = run._r.get_or_add_rPr()
                        rFonts = OxmlElement('w:rFonts')
                        rFonts.set(qn('w:eastAsia'), '微軟正黑體')
                        rFonts.set(qn('w:ascii'), '微軟正黑體')
                        rFonts.set(qn('w:hAnsi'), '微軟正黑體')
                        rPr.append(rFonts)

        for idx, (dom_name, dom_info) in enumerate(sorted_domains):
            r_idx = idx + 2
            row = table.rows[r_idx]
            q_list = dom_info['questions']
            q_str = '、'.join(map(str, sorted(q_list)))
            hi_avg = sum(questions_data[q]['hi'] for q in q_list if q in questions_data) / max(len(q_list), 1)
            all_avg = sum(questions_data[q]['all'] for q in q_list if q in questions_data) / max(len(q_list), 1)
            lo_avg = sum(questions_data[q]['lo'] for q in q_list if q in questions_data) / max(len(q_list), 1)
            set_cell_value(row.cells[0], dom_name)
            set_cell_value(row.cells[1], q_str)
            set_cell_value(row.cells[2], format_avg_rate(hi_avg))
            set_cell_value(row.cells[3], format_avg_rate(all_avg))
            set_cell_value(row.cells[4], format_avg_rate(lo_avg))

        overall_row = table.rows[num_domains + 2]
        set_cell_value(overall_row.cells[0], '整體', bold=True)
        total_q = sum(len(d['questions']) for _, d in sorted_domains)
        set_cell_value(overall_row.cells[1], str(total_q), bold=True)
        overall_hi = sum(q['hi'] for q in questions_data.values()) / max(len(questions_data), 1)
        overall_all = sum(q['all'] for q in questions_data.values()) / max(len(questions_data), 1)
        overall_lo = sum(q['lo'] for q in questions_data.values()) / max(len(questions_data), 1)
        set_cell_value(overall_row.cells[2], format_avg_rate(overall_hi), bold=True)
        set_cell_value(overall_row.cells[3], format_avg_rate(overall_all), bold=True)
        set_cell_value(overall_row.cells[4], format_avg_rate(overall_lo), bold=True)

        for row in table.rows:
            row.height = Pt(24)
            row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

        doc.save(output_path)

    elif subject_code == "M":
        content_domains = []
        for d1, d2 in domains:
            clean = str(d1).strip()
            if clean and clean not in content_domains:
                content_domains.append(clean)

        if grade_num in [3, 4, 5, 6] or is_entrant:
            order = ['數與計算', '量與實測', '空間與形狀', '關係', '資料與不確定性']
        else:
            order = ['數與量', '空間與形狀與坐標幾何', '代數與函數', '資料與不確定性']
        content_domains = sorted(content_domains, key=lambda x: order.index(x) if x in order else 999)
        cognitive_domains = ['概念理解', '程序執行', '解題思考']

        intersection = {(c, d): [] for c in content_domains for d in cognitive_domains}
        for num, (d1, d2) in enumerate(domains, start=1):
            c_val, d_val = str(d1).strip(), str(d2).strip()
            if (c_val, d_val) in intersection:
                intersection[(c_val, d_val)].append(num)

        doc = docx.Document(template_path)
        if len(doc.paragraphs) > 2:
            p2_text = f"由表6-1觀之，本次測驗全體學生的平均通過率是{overall_pass_pct}%。"
            set_paragraph_text_font(doc.paragraphs[2], p2_text)

        for p in doc.paragraphs:
            if re.match(r'^表\s*6-1', p.text.strip()):
                set_paragraph_text_font(p, "表6-1　整體學生及高分組、低分組學生於各評量向度之通過率")
                break

        table = doc.tables[0]
        math_header = "六" if is_entrant else grade_cn
        table.cell(0, 0).merge(table.cell(1, 0))
        set_cell_value(table.cell(0, 0), math_header, bold=True)

        for r_idx in [0, 1]:
            for c_idx, cell in enumerate(table.rows[r_idx].cells):
                fs = 10.5 if (r_idx == 1 and c_idx > 0) else 12
                set_cell_value(cell, cell.text, bold=True, font_size_pt=fs, template_cell=cell)

        src_tr = table.rows[2]._tr
        while len(table.rows) > 2:
            tr = table.rows[2]._tr
            table._tbl.remove(tr)

        for c_idx, c_name in enumerate(content_domains):
            start_row = 2 + 3 * c_idx
            for grp_idx, grp in enumerate(['hi', 'all', 'lo']):
                row = add_cloned_row(table, src_tr)
                grp_cn = {'hi': '高', 'all': '全', 'lo': '低'}[grp]
                set_cell_value(row.cells[0], format_domain_name_two_by_two(c_name))
                set_cell_value(row.cells[2], grp_cn)
                set_cell_value(row.cells[5], grp_cn)
                set_cell_value(row.cells[8], grp_cn)
                set_cell_value(row.cells[11], grp_cn)

            q_list_C = []
            for d_name in cognitive_domains:
                q_list_C.extend(intersection[(c_name, d_name)])

            table.cell(start_row, 0).merge(table.cell(start_row + 2, 0))
            set_cell_value(table.cell(start_row, 0), format_domain_name_two_by_two(c_name))
            table.cell(start_row, 10).merge(table.cell(start_row + 2, 10))
            set_cell_value(table.cell(start_row, 10), format_total_questions(len(q_list_C)))

            for grp_idx, grp in enumerate(['hi', 'all', 'lo']):
                r_idx = start_row + grp_idx
                if q_list_C:
                    avg_val = sum(questions_data[q][grp] for q in q_list_C if q in questions_data) / len(q_list_C)
                    set_cell_value(table.cell(r_idx, 12), format_avg_rate(avg_val))
                else:
                    set_cell_value(table.cell(r_idx, 12), '')
                if grp == 'all':
                    set_cell_shading(table.cell(r_idx, 11), "E0E0E0")
                    set_cell_shading(table.cell(r_idx, 12), "E0E0E0")

            for d_idx, d_name in enumerate(cognitive_domains):
                q_list = intersection[(c_name, d_name)]
                col_idx = 1 + 3 * d_idx
                if q_list:
                    q_str, align = format_question_list(q_list, style="math")
                    table.cell(start_row, col_idx).merge(table.cell(start_row + 2, col_idx))
                    set_cell_value(table.cell(start_row, col_idx), q_str, align=align)
                    for grp_idx, grp in enumerate(['hi', 'all', 'lo']):
                        r_idx = start_row + grp_idx
                        avg_val = sum(questions_data[q][grp] for q in q_list if q in questions_data) / len(q_list)
                        set_cell_value(table.cell(r_idx, col_idx + 2), format_avg_rate(avg_val))
                        if grp == 'all':
                            set_cell_shading(table.cell(r_idx, col_idx + 1), "E0E0E0")
                            set_cell_shading(table.cell(r_idx, col_idx + 2), "E0E0E0")
                else:
                    for r in range(start_row, start_row + 3):
                        table.cell(r, col_idx).merge(table.cell(r, col_idx + 2))
                    merged_cell = table.cell(start_row, col_idx).merge(table.cell(start_row + 2, col_idx))
                    set_cell_value(merged_cell, '')
                    set_cell_diagonal_border(merged_cell)

        summary_start_row = 2 + 3 * len(content_domains)
        for grp_idx, grp in enumerate(['hi', 'all', 'lo']):
            row = add_cloned_row(table, src_tr)
            grp_cn = {'hi': '高', 'all': '全', 'lo': '低'}[grp]
            set_cell_value(row.cells[0], '整體', bold=True)
            set_cell_value(row.cells[2], grp_cn, bold=True)
            set_cell_value(row.cells[5], grp_cn, bold=True)
            set_cell_value(row.cells[8], grp_cn, bold=True)
            set_cell_value(row.cells[11], grp_cn, bold=True)

        table.cell(summary_start_row, 0).merge(table.cell(summary_start_row + 2, 0))
        set_cell_value(table.cell(summary_start_row, 0), '整體', bold=True)
        total_q_count = len(questions_data)
        table.cell(summary_start_row, 10).merge(table.cell(summary_start_row + 2, 10))
        set_cell_value(table.cell(summary_start_row, 10), format_total_questions(total_q_count), bold=True)

        for grp_idx, grp in enumerate(['hi', 'all', 'lo']):
            r_idx = summary_start_row + grp_idx
            avg_val = sum(q[grp] for q in questions_data.values()) / max(len(questions_data), 1)
            set_cell_value(table.cell(r_idx, 12), format_avg_rate(avg_val), bold=True)
            if grp == 'all':
                set_cell_shading(table.cell(r_idx, 11), "E0E0E0")
                set_cell_shading(table.cell(r_idx, 12), "E0E0E0")

        for d_idx, d_name in enumerate(cognitive_domains):
            q_list_D = []
            for c_name in content_domains:
                q_list_D.extend(intersection[(c_name, d_name)])
            col_idx = 1 + 3 * d_idx
            table.cell(summary_start_row, col_idx).merge(table.cell(summary_start_row + 2, col_idx))
            set_cell_value(table.cell(summary_start_row, col_idx), format_total_questions(len(q_list_D)), bold=True)
            for grp_idx, grp in enumerate(['hi', 'all', 'lo']):
                r_idx = summary_start_row + grp_idx
                if q_list_D:
                    avg_val = sum(questions_data[q][grp] for q in q_list_D if q in questions_data) / len(q_list_D)
                    set_cell_value(table.cell(r_idx, col_idx + 2), format_avg_rate(avg_val), bold=True)
                else:
                    set_cell_value(table.cell(r_idx, col_idx + 2), '', bold=True)
                if grp == 'all':
                    set_cell_shading(table.cell(r_idx, col_idx + 1), "E0E0E0")
                    set_cell_shading(table.cell(r_idx, col_idx + 2), "E0E0E0")

        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            cantSplit = trPr.find(qn('w:cantSplit'))
            if cantSplit is None:
                cantSplit = OxmlElement('w:cantSplit')
                trPr.append(cantSplit)

        for r_idx in [0, 1]:
            trPr = table.rows[r_idx]._tr.get_or_add_trPr()
            tblHeader = trPr.find(qn('w:tblHeader'))
            if tblHeader is None:
                tblHeader = OxmlElement('w:tblHeader')
                trPr.append(tblHeader)

        col_widths = [
            Twips(847),
            Twips(1450), Twips(500), Twips(650),
            Twips(1150), Twips(500), Twips(650),
            Twips(1150), Twips(500), Twips(650),
            Twips(500), Twips(500), Twips(729)
        ]
        set_table_column_widths(table, col_widths)
        update_table_grid_widths(table, col_widths)
        set_table_cell_margins(table, top=60, bottom=60, left=45, right=45)
        doc.save(output_path)

    elif subject_code == "E":
        domain_questions = {}
        for num, (d1, d2) in enumerate(domains, start=1):
            m = re.match(r'^([AB]\d+)(.+)$', str(d2).strip())
            sub_name = m.group(2) if m else str(d2).strip()
            key = (str(d1).strip(), sub_name)
            if key not in domain_questions:
                domain_questions[key] = []
            domain_questions[key].append(num)

        sorted_keys = sorted(domain_questions.keys(), key=lambda x: min(domain_questions[x]))

        doc = docx.Document(template_path)
        if len(doc.paragraphs) > 1:
            doc.paragraphs[1].text = f"由表6-1觀之，本次測驗全體學生的平均通過率是{overall_pass_pct}%。"
            set_paragraph_text_font(doc.paragraphs[1], doc.paragraphs[1].text)

        for p in doc.paragraphs:
            if re.match(r'^表\s*6-1', p.text.strip()):
                set_paragraph_text_font(p, "表6-1　整體學生及高分組、低分組學生於各評量向度之通過率")
                break

        table = doc.tables[0]
        for r_idx in [0, 1]:
            row = table.rows[r_idx]
            row.height = Pt(24)
            row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                    p.paragraph_format.line_spacing = Pt(24)
                    for run in p.runs:
                        run.font.name = '微軟正黑體'
                        rPr = run._r.get_or_add_rPr()
                        rFonts = OxmlElement('w:rFonts')
                        rFonts.set(qn('w:eastAsia'), '微軟正黑體')
                        rFonts.set(qn('w:ascii'), '微軟正黑體')
                        rFonts.set(qn('w:hAnsi'), '微軟正黑體')
                        rPr.append(rFonts)

        while len(table.rows) > 2:
            tr = table.rows[2]._tr
            table._tbl.remove(tr)

        listen_rows = 0
        read_rows = 0
        for (m_val, s_name) in sorted_keys:
            q_list = domain_questions[(m_val, s_name)]
            q_str = '、'.join(map(str, sorted(q_list)))
            hi_avg = sum(questions_data[q]['hi'] for q in q_list if q in questions_data) / len(q_list)
            all_avg = sum(questions_data[q]['all'] for q in q_list if q in questions_data) / len(q_list)
            lo_avg = sum(questions_data[q]['lo'] for q in q_list if q in questions_data) / len(q_list)

            row = table.add_row()
            row.height = Pt(24)
            row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            set_cell_value(row.cells[0], m_val)
            set_cell_value(row.cells[1], s_name)
            set_cell_value(row.cells[2], q_str)
            set_cell_value(row.cells[3], format_avg_rate(hi_avg))
            set_cell_value(row.cells[4], format_avg_rate(all_avg))
            set_cell_value(row.cells[5], format_avg_rate(lo_avg))

            if m_val == '聽力':
                listen_rows += 1
            else:
                read_rows += 1

        if listen_rows > 1:
            merged_cell = table.cell(2, 0).merge(table.cell(2 + listen_rows - 1, 0))
            set_cell_value(merged_cell, '聽力')
        elif listen_rows == 1:
            set_cell_value(table.cell(2, 0), '聽力')

        if read_rows > 1:
            merged_cell = table.cell(2 + listen_rows, 0).merge(table.cell(2 + listen_rows + read_rows - 1, 0))
            set_cell_value(merged_cell, '閱讀')
        elif read_rows == 1:
            set_cell_value(table.cell(2 + listen_rows, 0), '閱讀')

        row = table.add_row()
        row.height = Pt(24)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        set_cell_value(row.cells[0], '整體', bold=True)
        set_cell_value(row.cells[1], '整體', bold=True)
        set_cell_value(row.cells[2], '', bold=True)

        overall_hi = sum(q['hi'] for q in questions_data.values()) / max(len(questions_data), 1)
        overall_all = sum(q['all'] for q in questions_data.values()) / max(len(questions_data), 1)
        overall_lo = sum(q['lo'] for q in questions_data.values()) / max(len(questions_data), 1)
        set_cell_value(row.cells[3], format_avg_rate(overall_hi), bold=True)
        set_cell_value(row.cells[4], format_avg_rate(overall_all), bold=True)
        set_cell_value(row.cells[5], format_avg_rate(overall_lo), bold=True)

        total_rows = len(table.rows)
        merged_cell = table.cell(total_rows - 1, 0).merge(table.cell(total_rows - 1, 1))
        set_cell_value(merged_cell, '整體', bold=True)
        doc.save(output_path)

    elif subject_code == "S":
        content_domains = []
        for d1, _ in domains:
            m = re.search(r"^(.+?)[（(]", str(d1))
            name = m.group(1).strip() if m else str(d1).strip()
            if name and name not in content_domains:
                content_domains.append(name)

        cognitive_domains = ['記憶', '理解', '應用', '分析']
        intersection = {(c, d): [] for c in content_domains for d in cognitive_domains}

        def parse_sci_cog(raw):
            m = re.match(r"([A-D])\s*(.+)", str(raw))
            return m.group(2).strip() if m else str(raw).strip()

        for num, (d1, d2) in enumerate(domains, start=1):
            m = re.search(r"^(.+?)[（(]", str(d1))
            c_val = m.group(1).strip() if m else str(d1).strip()
            d_val = parse_sci_cog(d2)
            if (c_val, d_val) in intersection:
                intersection[(c_val, d_val)].append(num)

        doc = docx.Document(template_path)
        if len(doc.paragraphs) > 2:
            p2_text = f"由表6-1觀之，本次測驗全體學生的平均通過率是{overall_pass_pct}%。"
            set_paragraph_text_font(doc.paragraphs[2], p2_text)

        for p in doc.paragraphs:
            if re.match(r'^表\s*6-1', p.text.strip()):
                set_paragraph_text_font(p, "表6-1　整體學生及高分組、低分組學生於各評量向度之通過率")
                break

        old_table = doc.tables[0]
        num_rows = 2 + 3 * len(content_domains) + 3
        new_table = doc.add_table(rows=num_rows, cols=16)

        new_table._tbl.tblPr.clear()
        for child in old_table._tbl.tblPr:
            new_table._tbl.tblPr.append(copy.deepcopy(child))

        for row in new_table.rows:
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                vAlign = tcPr.find(qn('w:vAlign'))
                if vAlign is None:
                    vAlign = OxmlElement('w:vAlign')
                    vAlign.set(qn('w:val'), 'center')
                    tcPr.append(vAlign)

        for r_idx in [0, 1]:
            for cell in new_table.rows[r_idx].cells:
                set_cell_shading(cell, "E0E0E0")

        def set_cell_value_S(r, c, text, bold=False, font_size_pt=12, align=WD_ALIGN_PARAGRAPH.CENTER):
            cell = new_table.cell(r, c)
            set_cell_value(cell, text, bold=bold, align=align, font_size_pt=font_size_pt)

        sci_header = "六" if is_entrant else grade_cn
        # Row 0
        row0_vals = [sci_header] + ['記憶']*3 + ['理解']*3 + ['應用']*3 + ['分析']*3 + ['整體']*3
        for i, val in enumerate(row0_vals):
            set_cell_value_S(0, i, val, bold=True, font_size_pt=12)

        # Row 1
        row1_vals = [sci_header] + ['題號', '分組', '通過率']*4 + ['題數', '分組', '通過率']
        for i, val in enumerate(row1_vals):
            fs = 10.5 if i > 0 else 12
            set_cell_value_S(1, i, val, bold=True, font_size_pt=fs)

        new_table.cell(0, 0).merge(new_table.cell(1, 0))
        set_cell_value_S(0, 0, sci_header, bold=True, font_size_pt=12)

        new_table.cell(0, 1).merge(new_table.cell(0, 3))
        set_cell_value_S(0, 1, '記憶', bold=True, font_size_pt=12)
        new_table.cell(0, 4).merge(new_table.cell(0, 6))
        set_cell_value_S(0, 4, '理解', bold=True, font_size_pt=12)
        new_table.cell(0, 7).merge(new_table.cell(0, 9))
        set_cell_value_S(0, 7, '應用', bold=True, font_size_pt=12)
        new_table.cell(0, 10).merge(new_table.cell(0, 12))
        set_cell_value_S(0, 10, '分析', bold=True, font_size_pt=12)
        new_table.cell(0, 13).merge(new_table.cell(0, 15))
        set_cell_value_S(0, 13, '整體', bold=True, font_size_pt=12)

        for c_idx, c_name in enumerate(content_domains):
            start_row = 2 + 3 * c_idx
            for grp_idx, grp in enumerate(['hi', 'all', 'lo']):
                r_idx = start_row + grp_idx
                grp_cn = {'hi': '高', 'all': '全', 'lo': '低'}[grp]
                set_cell_value_S(r_idx, 0, format_domain_name_two_by_two(c_name))
                set_cell_value_S(r_idx, 2, grp_cn)
                set_cell_value_S(r_idx, 5, grp_cn)
                set_cell_value_S(r_idx, 8, grp_cn)
                set_cell_value_S(r_idx, 11, grp_cn)
                set_cell_value_S(r_idx, 14, grp_cn)

            q_list_C = []
            for d_name in cognitive_domains:
                q_list_C.extend(intersection[(c_name, d_name)])

            new_table.cell(start_row, 0).merge(new_table.cell(start_row + 2, 0))
            set_cell_value_S(start_row, 0, format_domain_name_two_by_two(c_name))
            new_table.cell(start_row, 13).merge(new_table.cell(start_row + 2, 13))
            set_cell_value_S(start_row, 13, format_total_questions(len(q_list_C)))

            for grp_idx, grp in enumerate(['hi', 'all', 'lo']):
                r_idx = start_row + grp_idx
                if q_list_C:
                    avg_val = sum(questions_data[q][grp] for q in q_list_C if q in questions_data) / len(q_list_C)
                    set_cell_value_S(r_idx, 15, format_avg_rate(avg_val))
                else:
                    set_cell_value_S(r_idx, 15, '')
                if grp == 'all':
                    set_cell_shading(new_table.cell(r_idx, 14), "E0E0E0")
                    set_cell_shading(new_table.cell(r_idx, 15), "E0E0E0")

            for d_idx, d_name in enumerate(cognitive_domains):
                q_list = intersection[(c_name, d_name)]
                col_idx = 1 + 3 * d_idx
                if q_list:
                    q_str, align = format_question_list(q_list, style="science")
                    new_table.cell(start_row, col_idx).merge(new_table.cell(start_row + 2, col_idx))
                    set_cell_value_S(start_row, col_idx, q_str, align=align)
                    for grp_idx, grp in enumerate(['hi', 'all', 'lo']):
                        r_idx = start_row + grp_idx
                        avg_val = sum(questions_data[q][grp] for q in q_list if q in questions_data) / len(q_list)
                        set_cell_value_S(r_idx, col_idx + 2, format_avg_rate(avg_val))
                        if grp == 'all':
                            set_cell_shading(new_table.cell(r_idx, col_idx + 1), "E0E0E0")
                            set_cell_shading(new_table.cell(r_idx, col_idx + 2), "E0E0E0")
                else:
                    for r in range(start_row, start_row + 3):
                        new_table.cell(r, col_idx).merge(new_table.cell(r, col_idx + 2))
                    merged_cell = new_table.cell(start_row, col_idx).merge(new_table.cell(start_row + 2, col_idx))
                    set_cell_value(merged_cell, '')
                    set_cell_diagonal_border(merged_cell)

        summary_start_row = 2 + 3 * len(content_domains)
        for grp_idx, grp in enumerate(['hi', 'all', 'lo']):
            r_idx = summary_start_row + grp_idx
            grp_cn = {'hi': '高', 'all': '全', 'lo': '低'}[grp]
            set_cell_value_S(r_idx, 0, format_domain_name_two_by_two('整體'), bold=True)
            set_cell_value_S(r_idx, 2, grp_cn, bold=True)
            set_cell_value_S(r_idx, 5, grp_cn, bold=True)
            set_cell_value_S(r_idx, 8, grp_cn, bold=True)
            set_cell_value_S(r_idx, 11, grp_cn, bold=True)
            set_cell_value_S(r_idx, 14, grp_cn, bold=True)

        new_table.cell(summary_start_row, 0).merge(new_table.cell(summary_start_row + 2, 0))
        set_cell_value_S(summary_start_row, 0, format_domain_name_two_by_two('整體'), bold=True)
        total_q_count = len(questions_data)
        new_table.cell(summary_start_row, 13).merge(new_table.cell(summary_start_row + 2, 13))
        set_cell_value_S(summary_start_row, 13, format_total_questions(total_q_count), bold=True)

        for grp_idx, grp in enumerate(['hi', 'all', 'lo']):
            r_idx = summary_start_row + grp_idx
            avg_val = sum(q[grp] for q in questions_data.values()) / max(len(questions_data), 1)
            set_cell_value_S(r_idx, 15, format_avg_rate(avg_val), bold=True)
            if grp == 'all':
                set_cell_shading(new_table.cell(r_idx, 14), "E0E0E0")
                set_cell_shading(new_table.cell(r_idx, 15), "E0E0E0")

        for d_idx, d_name in enumerate(cognitive_domains):
            q_list_D = []
            for c_name in content_domains:
                q_list_D.extend(intersection[(c_name, d_name)])
            col_idx = 1 + 3 * d_idx
            new_table.cell(summary_start_row, col_idx).merge(new_table.cell(summary_start_row + 2, col_idx))
            set_cell_value_S(summary_start_row, col_idx, format_total_questions(len(q_list_D)), bold=True)
            for grp_idx, grp in enumerate(['hi', 'all', 'lo']):
                r_idx = summary_start_row + grp_idx
                if q_list_D:
                    avg_val = sum(questions_data[q][grp] for q in q_list_D if q in questions_data) / len(q_list_D)
                    set_cell_value_S(r_idx, col_idx + 2, format_avg_rate(avg_val), bold=True)
                else:
                    set_cell_value_S(r_idx, col_idx + 2, '', bold=True)
                if grp == 'all':
                    set_cell_shading(new_table.cell(r_idx, col_idx + 1), "E0E0E0")
                    set_cell_shading(new_table.cell(r_idx, col_idx + 2), "E0E0E0")

        for r_idx in [0, 1]:
            trPr = new_table.rows[r_idx]._tr.get_or_add_trPr()
            tblHeader = trPr.find(qn('w:tblHeader'))
            if tblHeader is None:
                tblHeader = OxmlElement('w:tblHeader')
                trPr.append(tblHeader)

        for row in new_table.rows:
            trPr = row._tr.get_or_add_trPr()
            cantSplit = trPr.find(qn('w:cantSplit'))
            if cantSplit is None:
                cantSplit = OxmlElement('w:cantSplit')
                trPr.append(cantSplit)

        col_widths = [
            Twips(600),
            Twips(550), Twips(450), Twips(835),
            Twips(550), Twips(450), Twips(835),
            Twips(550), Twips(450), Twips(835),
            Twips(550), Twips(450), Twips(835),
            Twips(550), Twips(450), Twips(836)
        ]
        set_table_column_widths(new_table, col_widths)
        update_table_grid_widths(new_table, col_widths)
        set_table_cell_margins(new_table, top=60, bottom=60, left=30, right=30)

        old_table._tbl.addprevious(new_table._tbl)
        old_table._tbl.getparent().remove(old_table._tbl)
        doc.save(output_path)

    from report_images import normalize_chapter6_layout
    normalize_chapter6_layout(output_path, subject_code)
    if analysis.get('draft'):
        from report_images import mark_missing_statistics
        mark_missing_statistics(output_path, 6)
    from proper_names import format_counties
    format_counties(output_path)
    print(f" -> 產出成功: {output_path}")

# ==============================================================================
# PIPELINE COORDINATOR
# ==============================================================================
def resolve_templates(template_dir, subject):
    sub = subject.upper()[0]
    t4 = None
    t6 = None

    if sub == "C":
        t4 = os.path.join(template_dir, "pydocxTemplet_C_new.docx")
        if not os.path.exists(t4):
            t4 = os.path.join(template_dir, "pydocxTemplet_C.docx")
        t6 = os.path.join(template_dir, "C_demo.docx")
    elif sub == "M":
        t4 = os.path.join(template_dir, "pydocxTemplet_M.docx")
        t6 = os.path.join(template_dir, "S_M_demo.docx")
    elif sub == "E":
        t4 = os.path.join(template_dir, "pydocxTemplet_E.docx")
        t6 = os.path.join(template_dir, "six_templet_E.docx")
    elif sub == "S":
        t4 = os.path.join(template_dir, "pydocxTemplet_S.docx")
        t6 = os.path.join(template_dir, "S_M_demo.docx")

    return t4, t6

def process_single_run(ans_file, analysis_file, subject, grade=None, year="115", chapter="both", template_dir=None, out_dir=None, images_manifest=None, booklet_pdf=None, output_suffix=''):
    if not os.path.exists(ans_file):
        raise FileNotFoundError(f"找不到答案檔: {ans_file}")
    if not os.path.exists(analysis_file):
        raise FileNotFoundError(f"找不到試題分析總表: {analysis_file}")

    if template_dir is None:
        script_dir = os.path.dirname(os.path.abspath(__file__))
        template_dir = os.path.join(os.path.dirname(script_dir), "resources", "templates")
        if not os.path.exists(template_dir):
            template_dir = os.getcwd()

    if out_dir is None:
        out_dir = os.getcwd()
    os.makedirs(out_dir, exist_ok=True)
    if images_manifest and booklet_pdf:
        raise ValueError('booklet_pdf 與 images_manifest 請擇一')
    if booklet_pdf:
        from question_screenshots import export, identity
        from pathlib import Path
        crop_output = Path(out_dir)/'screenshots'
        export(Path(booklet_pdf), 'booklets', crop_output, 220, {}, {})
        images_manifest = crop_output/'booklets'/identity(Path(booklet_pdf))/'manifest.json'

    subject_code = subject.upper()[0]
    t4_path, t6_path = resolve_templates(template_dir, subject_code)

    grades_data = read_answers_and_domains(ans_file, subject_code)
    analysis_data = read_analysis_file(analysis_file)

    target_grades = []
    if grade and str(grade).lower() != "all":
        target_grades = [str(grade)]
    else:
        # Match intersection of grades between answers and analysis
        common_keys = set(grades_data.keys()) & set(analysis_data.keys())
        if not common_keys:
            # Fallback: try numeric string match
            for gk in grades_data:
                for ak in analysis_data:
                    if str(gk) == str(ak) or re.sub(r'\D', '', str(gk)) == re.sub(r'\D', '', str(ak)):
                        common_keys.add(gk)
        if not common_keys:
            # If both single
            if len(grades_data) == 1 and len(analysis_data) == 1:
                target_grades = [list(grades_data.keys())[0]]
            else:
                target_grades = list(grades_data.keys())
        else:
            target_grades = sorted(list(common_keys), key=lambda x: str(x))

    print(f"\n=======================================================")
    print(f"開始處理: 科目={SUBJECT_NAME_MAP.get(subject_code, subject_code)} ({subject_code}), 年級={target_grades}")
    print(f"答案檔: {os.path.basename(ans_file)}")
    print(f"分析總表: {os.path.basename(analysis_file)}")
    print(f"輸出目錄: {out_dir}")
    print(f"=======================================================")

    if images_manifest and len(target_grades) != 1:
        raise ValueError('單一題本／圖片清單只可套用單一年級，請明確指定 --grade')
    for g in target_grades:
        g_data = grades_data.get(g) or list(grades_data.values())[0]
        a_data = analysis_data.get(g) or analysis_data.get(str(g)) or list(analysis_data.values())[0]

        answers = g_data["answers"]
        domains = g_data["domains"]

        g_tag = str(g)
        if g_tag == "single":
            # Extract from filename
            m = re.search(r'([A-Za-z]\d+N?)', os.path.basename(ans_file))
            g_tag = m.group(1) if m else "7N"

        cohort_code = f"{subject_code}{g_tag}" if not g_tag.startswith(subject_code) else g_tag

        # Chapter 4
        if chapter in ["4", "both"]:
            if t4_path and os.path.exists(t4_path):
                out_name_c4 = f"第四章_{cohort_code}{output_suffix}.docx"
                out_path_c4 = os.path.join(out_dir, out_name_c4)
                generate_chapter_4(subject_code, g_tag, year, answers, domains, a_data, t4_path, out_path_c4, images_manifest)
            else:
                print(f"[警告] 找不到第四章樣板: {t4_path}")

        # Chapter 6
        if chapter in ["6", "both"]:
            if t6_path and os.path.exists(t6_path):
                out_name_c6 = f"第六章_{cohort_code}{output_suffix}.docx"
                out_path_c6 = os.path.join(out_dir, out_name_c6)
                generate_chapter_6(subject_code, g_tag, year, answers, domains, a_data, t6_path, out_path_c6)
            else:
                print(f"[警告] 找不到第六章樣板: {t6_path}")

def auto_scan_and_run(scan_dir, out_dir=None, year="115", chapter="both", output_suffix=''):
    if out_dir is None:
        out_dir = scan_dir
    files = os.listdir(scan_dir)
    print(f"自動掃描目錄: {scan_dir}")

    # Map by subject / cohort
    # Pattern 1: e.g. C7N_ans.xlsx + C7N_試題分析總表.xlsx
    # Pattern 2: e.g. 115_C_ans.xlsx + 115_C_試題分析總表.xlsx
    ans_files = [f for f in files if "ans" in f.lower() and f.endswith(".xlsx") and not f.startswith("~$")]
    anal_files = [f for f in files if "分析總表" in f and f.endswith(".xlsx") and not f.startswith("~$")]

    subjects = ["C", "M", "E", "S"]
    for sub in subjects:
        # Find matching pairs
        sub_ans = [f for f in ans_files if f"_{sub}_" in f or f.startswith(f"{sub}") or f"_{sub}" in f]
        sub_anal = [f for f in anal_files if f"_{sub}_" in f or f.startswith(f"{sub}") or f"_{sub}" in f]

        for a_file in sub_ans:
            # Match analysis file
            base_key = re.sub(r'_ans.*', '', a_file, flags=re.I)
            match_anal = None
            for an in sub_anal:
                if base_key in an or re.sub(r'_試題分析.*', '', an) in a_file:
                    match_anal = an
                    break
            if match_anal is None and len(sub_anal) == 1:
                match_anal = sub_anal[0]

            if match_anal:
                p_ans = os.path.join(scan_dir, a_file)
                p_anal = os.path.join(scan_dir, match_anal)
                try:
                    process_single_run(p_ans, p_anal, subject=sub, year=year, chapter=chapter, out_dir=out_dir, output_suffix=output_suffix)
                except Exception as e:
                    print(f"[錯誤] 處理 {a_file} 失敗: {e}")

def main():
    parser = argparse.ArgumentParser(description="SAAA Chapter 4 & 6 Report Generator")
    parser.add_argument("--ans-file", help="Path to answers/dimensions Excel file")
    parser.add_argument("--analysis-file", help="Path to item analysis Excel file")
    parser.add_argument("--subject", choices=["C", "M", "E", "S", "all"], help="Subject code (C: Chinese, M: Math, E: English, S: Science)")
    parser.add_argument("--grade", default="all", help="Grade number (3..8, 7N, or all)")
    parser.add_argument("--year", default="115", help="School year string (e.g., 115)")
    parser.add_argument("--chapter", choices=["4", "6", "both"], default="both", help="Chapter to generate (4, 6, or both)")
    parser.add_argument("--template-dir", help="Directory containing template docx files")
    parser.add_argument("--out-dir", help="Directory to save generated docx reports")
    parser.add_argument("--auto-scan", help="Scan a directory and automatically run for all found datasets")
    parser.add_argument('--booklet-pdf', help='題本 PDF，自動截圖並填入第四章')
    parser.add_argument('--images-manifest', help='既有截圖 manifest.json，與 --booklet-pdf 擇一')
    parser.add_argument('--output-suffix', default='', help='輸出檔名後綴，例如 _AI產出')

    args = parser.parse_args()

    if args.auto_scan:
        if args.images_manifest or args.booklet_pdf:
            parser.error('--auto-scan 的多科檔案不能共用同一份題本；請使用 build_from_source.py 或單科參數')
        auto_scan_and_run(args.auto_scan, out_dir=args.out_dir, year=args.year, chapter=args.chapter, output_suffix=args.output_suffix)
        return

    if not args.ans_file or not args.analysis_file or not args.subject:
        parser.print_help()
        sys.exit(1)

    process_single_run(
        ans_file=args.ans_file,
        analysis_file=args.analysis_file,
        subject=args.subject,
        grade=args.grade,
        year=args.year,
        chapter=args.chapter,
        template_dir=args.template_dir,
        out_dir=args.out_dir,
        images_manifest=args.images_manifest,
        booklet_pdf=args.booklet_pdf,
        output_suffix=args.output_suffix,
    )

if __name__ == "__main__":
    main()
