"""唯讀量測指定四組參考 Word；將物理尺寸與 XML 明確設定分開呈現。"""
import argparse
from collections import Counter
import html
import json
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

GROUPS = {
    '國語文': '國語文/提供領域_第4、6章數據',
    '英語文（國小）': '英語文(國小)/01 內部初稿',
    '英語文（國中）': '英語文(國中)/01 內部初稿',
    '自然科': '自然/01_原稿_提供領域版',
}


def cm(value):
    return round(int(value)/360000, 3) if value is not None else None


def measure(path, root):
    doc = Document(path)
    sections = []
    for sec in doc.sections:
        sections.append(dict(page_cm=[cm(sec.page_width), cm(sec.page_height)],
            margins_cm=[cm(sec.top_margin), cm(sec.bottom_margin), cm(sec.left_margin), cm(sec.right_margin)],
            usable_cm=[cm(sec.page_width-sec.left_margin-sec.right_margin),
                       cm(sec.page_height-sec.top_margin-sec.bottom_margin)]))
    grids = []
    for table in doc.tables:
        widths = [cm(col.width) for col in table.columns]
        if widths not in grids:
            grids.append(widths)
    # 不用 row.cells 重複計算水平／垂直合併儲存格。
    alignment, spacing, sizes = Counter(), Counter(), Counter()
    for tc in doc.element.body.iter(qn('w:tc')):
        props = tc.find(qn('w:tcPr'))
        vertical = props.find(qn('w:vAlign')) if props is not None else None
        alignment[vertical.get(qn('w:val')) if vertical is not None else '未明設（繼承／預設）'] += 1
    for para in doc.element.body.iter(qn('w:p')):
        text = ''.join(para.itertext())
        if not text.strip():
            continue
        props = para.find(qn('w:pPr'))
        node = props.find(qn('w:spacing')) if props is not None else None
        if node is not None and node.get(qn('w:line')):
            rule = node.get(qn('w:lineRule'), 'auto')
            value = int(node.get(qn('w:line')))
            label = f'{value/20:g} pt（{rule}）' if rule != 'auto' else f'{value/240:g} 倍（auto）'
        else:
            label = '未明設（繼承樣式／格線）'
        spacing[label] += 1
    for node in doc.element.body.iter(qn('w:sz')):
        sizes[str(int(node.get(qn('w:val')))/2)] += 1
    rows = []
    for table in doc.tables:
        if any('試題代碼' in c.text for row in table.rows[:2] for c in row.cells):
            rows = [dict(height_pt=round(r.height.pt, 2) if r.height else None,
                         rule=str(r.height_rule)) for r in table.rows[:9]]
            break
    return dict(file=str(path.relative_to(root)), sections=sections, table_grids_cm=grids,
                first_item_rows=rows, vertical_alignment=dict(alignment),
                paragraph_spacing=dict(spacing), explicit_run_font_sizes_pt=dict(sizes),
                images_cm=[[cm(s.width), cm(s.height)] for s in doc.inline_shapes])


def write_report(root, out):
    records = []
    panels = []
    for group, subdir in GROUPS.items():
        paths = sorted(p for p in (root/subdir).rglob('*.docx') if not p.name.startswith(('._', '~$')))
        current = [measure(p, root) for p in paths]
        records.extend(dict(group=group, **r) for r in current)
        rows = []
        for r in current:
            layouts = list({json.dumps(s, ensure_ascii=False): None for s in r['sections']})
            layout_text = '<br>'.join(html.escape(x) for x in layouts)
            widths = sorted({round(sum(g), 2) for g in r['table_grids_cm']})
            detail = {k:r[k] for k in ['first_item_rows', 'vertical_alignment', 'paragraph_spacing', 'explicit_run_font_sizes_pt', 'images_cm']}
            rows.append(f'<tr><td>{html.escape(r["file"])}</td><td>{layout_text}</td><td>{widths}</td><td>{len(r["images_cm"])}</td><td><details><summary>列高／對齊／行距／字級／圖片尺寸</summary><pre>{html.escape(json.dumps(detail, ensure_ascii=False, indent=2))}</pre></details></td></tr>')
        fonts, spacing = Counter(), Counter()
        for record in current:
            fonts.update(record['explicit_run_font_sizes_pt'])
            spacing.update(record['paragraph_spacing'])
        summary = html.escape(json.dumps(dict(explicit_run_font_sizes_pt=dict(fonts), paragraph_spacing=dict(spacing)), ensure_ascii=False, indent=2))
        panels.append(f'<h2>{group}：{len(current)} 份</h2><details><summary>本組字級／行距的明設數量統計</summary><pre>{summary}</pre></details><table><tr><th>檔案</th><th>各節版面（cm；邊界依序上、下、左、右）</th><th>表格格線總寬（cm）</th><th>內嵌圖數</th><th>詳細數字</th></tr>{"".join(rows)}</table>')
    out.mkdir(parents=True, exist_ok=True)
    (out/'原稿布局量測_AI產出.json').write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding='utf-8')
    intro = '''<h1>115 Word 原稿布局量測與套用決策</h1>
<p>唯讀分析使用者指定的四個資料夾，共32份 Word。所有尺寸取自文件設定，未修改原稿。頁面／邊界／表格／圖片以公分表示；字級與固定行高以點表示。</p>
<h2>國語文</h2><p>12份第4、6章原稿的首節皆為 A4（21×29.7），上下邊界2.54，左2.00、右1.75，可用17.25×24.62。代表題表格線總寬約17.23；格線是合併儲存格的底層結構，並非9個可見欄。跨來源頁的文章先合成整篇，再等比例放入同一報告頁；不將分段拆成多頁。題圖沿用約11點等效字級。</p>
<h2>英語文：國小與國中分開處理</h2><p>國小3份、國中2份原稿，首節同為 A4、可用約17.25×24.62。國小代表題表寬17.64，比首節可用頁寬多0.39；國中約17.23。後續節另見國小右邊界約1.50、可用寬17.50；國中右邊界2.00、可用寬17.00，不能只讀第一節，也不可把國小欄寬直接複製到新報告。題圖維持13點等效字級；共用閱讀文章按一頁可用寬高個別等比例縮放。只有題組編號的聽力材料保留在題目欄，不放大占整頁。</p>
<h2>自然科</h2><p>15份原稿包含原稿、修訂版與插圖版。首節均為A4；11份左右邊界約2.00／1.75，可用寬17.25；另4份左右各3.175，可用寬約14.65，不能套用同一組寬度。代表題表約17.23；各節及各表實際數字如下。原稿與修訂稿的題碼／向度曾有差異，版式參照不代表可取代正式答案及本次統計。</p>
<h2>本次生成規則</h2><ul><li>一般題圖靠左、靠上；國語、英文整頁共用文章在單頁文章區域垂直置中、水平靠左；保留圖片長寬比。</li><li>共用文章頁可用17.25×24.62；文章圖上限16.85×22.82，另留標題與分頁空間。先合併全部來源段落，再一次插入。</li><li>文章為符合整頁而使用獨立縮放倍率；一般題圖仍維持各科共同倍率。</li><li>分析文字固定24點行高；選答率與試題分析開頭同頁。去題號、去考場提示及取消黑色方點的既有規則保留。</li><li>桃園新生現有完整輸入為國語30題、英文35題、數學25題；本次驗證6份第4、6章報告。自然科僅分析參考稿，未冒用其他施測資料產出自然科報告。</li></ul>
<p>「未明設」不等於零或沒有格式，可能由樣式、文件預設或稿紙格線繼承。行距統計是正文所有非空段落的XML明設值，不等同試題分析區專屬行距；明設字級也不是圖片內字級。表格總寬是底層格線相加，須配合該節實際可用頁寬判讀。詳情保留完整數字供追溯。</p>'''
    (out/'各科布局分析_AI產出.html').write_text('<!doctype html><meta charset="utf-8"><title>各科布局分析</title><style>body{font:17px system-ui;line-height:1.7;margin:36px;color:#172b40}table{border-collapse:collapse;width:100%;font-size:14px}td,th{border:1px solid #ccd5de;padding:9px;text-align:left;vertical-align:top}pre{white-space:pre-wrap;max-height:500px;overflow:auto}h2{margin-top:34px}td:first-child{max-width:300px;overflow-wrap:anywhere}</style>'+intro+''.join(panels), encoding='utf-8')
    print(json.dumps({'documents':len(records), 'output':str(out)}, ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--dir', type=Path, required=True)
    parser.add_argument('--out-dir', type=Path, required=True)
    args = parser.parse_args()
    write_report(args.dir, args.out_dir)
